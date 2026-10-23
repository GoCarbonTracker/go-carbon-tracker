"""The Privacy-in-exports gate, 7.1's part (story 7.1 section G (a) and (c)): on tests/report_fixture.py, the
sanitised dataset holds no seed, the complete dataset holds exactly the seeds its fields can reach, credentials and
paths are in neither, the preview's counts add up, and sanitised(W) equals sanitised(W') once ids are renumbered in
first-appearance order (K12). The fixture's environment (GCT_MODELS and the api_key_env variable holding CRED-ENV-01)
is set for every build and every output, and store.now is fixed so W and W' write the same times.

Part B (story 7.7): on tests/leakage_fixture.py, every report format in both modes through gct report all and both
packages through gct package create, scanned from the files alone by tests/leakage.py: the seed matrix with the
never-list and its source proof (AC 7, 8), crop and marker hits per file, the pages in a browser (AC 11), the
injected row per renderer (AC 9), the packages' files (AC 10), the differential on every rendered format with its
named clock parts (AC 7a), section 4's behaviour (AC 12) and the coverage map (AC 13). Cited, not repeated (the gate
run includes them): tests/test_export_json.py::test_every_export_validates (schema validation, 7.2);
tests/test_render_pdf.py::test_found_recovers_a_planted_crop and
::test_a_no_box_ref_is_a_whole_page_entry_and_every_page_entry_is_the_viewer_render (readable, same-version crops,
7.4); tests/test_package.py::test_a_sanitised_package_holds_what_its_preview_says_and_reopens,
::test_a_complete_package_needs_include_private_and_carries_every_row, ::test_only_package_py_holds_the_phrase,
::test_no_command_calls_a_copied_or_changed_folder_sanitised (FR72),
::test_an_archived_shareable_document_ships_still_archived (7.6-P3), ::test_an_empty_sanitised_plan_previews_the_relabel_line_and_create_refuses (7.6-P8) and
::test_version_1_equals_the_source_version_for_shipped_rows (the package's own meta, 7.6)."""
import contextlib
import copy
import datetime
import gc
import hashlib
import io
import json
import os
import re
import shutil
import sqlite3
import tempfile
import time
import types
import zipfile
from collections import Counter
from pathlib import Path

import pytest
from PIL import Image

import leakage
import leakage_fixture as lf
import report_fixture as f
from gate_fixture import gct
from gct import cli
from gct.models import redact
from gct.reports import crops, dataset, export
from gct.workspace import layout, store

NOW = "2026-10-23T00:00:00+00:00"
FILTERS = [{}, {"review": ["disagreement"]}, {"review": ["needs-re-review"]}, {"comparison": ["unresolved"]},
           {"comparison": ["restatement"]}]


@pytest.fixture(scope="module")
def env(tmp_path_factory):
    """Ruling 1 (CRED-ENV-01): the fixture's models file and key variable, as tests/test_report_fixture.py sets them,
    for the whole module (every build and every output)."""
    models = tmp_path_factory.mktemp("models") / "models.toml"
    models.write_text(f.MODELS)
    with pytest.MonkeyPatch.context() as mp:
        mp.setenv("GCT_MODELS", str(models))
        mp.setenv(f.API_KEY_ENV, f.CRED_ENV)
        mp.setattr(store, "now", lambda: NOW)
        redact._names.cache_clear()
        assert redact._names() == (f.API_KEY_ENV,)
        yield
    redact._names.cache_clear()


@pytest.fixture(scope="module")
def w(env, tmp_path_factory):
    return f.build(tmp_path_factory.mktemp("w") / "ws")


@pytest.fixture(scope="module")
def w2(env, tmp_path_factory):
    return f.build(tmp_path_factory.mktemp("w2") / "ws", without_private=True)


def build(root, **kw):
    conn = store.connect(root / layout.DB_NAME)
    try:
        return dataset.build(conn, root, **kw)
    finally:
        conn.close()


def text(d) -> str:
    return json.dumps(d, ensure_ascii=False)


def found(d) -> set:
    """The seeds a dataset holds: a text seed or a hex digest anywhere in its JSON. Image seeds match pixels only,
    and a dataset holds none (C3d)."""
    t = text(d)
    return {s["name"] for s in f.SEEDS if s["kind"] != "image" and s["value"] in t}


# (a): the scan

@pytest.mark.parametrize("include", [(), ("contributor_text", "flags", "identity")])
def test_the_sanitised_dataset_holds_no_seed(w, env, include):
    for flt in FILTERS:
        d, p = build(w, filter=flt, include=include)
        hits = found(d)
        if include:  # the classes restore contributor text and names, never document text, credentials or paths
            hits -= {s["name"] for s in f.SEEDS if s["kind"] == "contributor"}
        assert hits == set(), (flt, hits)
        assert p["history_mark"] not in text(d)  # K2 (labels_mark is a small number; data_version has neither key)
        assert set(d["data_version"]) == {"version", "built_at", "sha256", "projection_schema"}


def test_the_complete_dataset_holds_exactly_its_reachable_seeds(w, env):
    d, _ = build(w, mode="complete")
    shipped = {e["sha256"] for e in d["evidence"]}
    unreachable = {"CON-ACCEPT-01", "DOC-BASIS-01"}  # no dataset field carries them (section G (a))
    want = {s["name"] for s in f.SEEDS if s["expect"] == "complete_only" and s["kind"] in ("document_text", "contributor")
            and s["name"] not in unreachable}
    want |= {s["name"] for s in f.SEEDS if s["kind"] == "file_hash" and s["value"] in shipped}
    assert found(d) == want
    # the scan can fail: each expected kind is there, and the archived document has no evidence entry
    assert {"FILE-PDF-01", "FILE-CSV-01", "FILE-PDF-03", "FILE-PDF-04"} <= want and "FILE-PDF-02" not in want
    assert {"DOC-NAME-02", "DOC-COMPANY-01", "CON-NAME-01", "DOC-REASON-01", "DOC-CELL-02"} <= want


@pytest.mark.parametrize("mode", dataset.MODES)
def test_credentials_and_paths_are_in_neither_mode(w, env, mode):
    for flt in FILTERS:
        d, p = build(w, mode=mode, filter=flt, include=dataset.INCLUDE)
        never = {s["name"] for s in f.SEEDS if s["expect"] == "never"}
        assert not {s["name"] for s in f.SEEDS if s["name"] in never and s["value"] in text(d) + text(p)}, flt
        assert str(w) not in text(d) and "/private/" not in text(d) and "/tmp" not in text(d)


def test_the_preview_counts_match_the_hand_counts(w, env):
    for flt in FILTERS:
        assert set(build(w, mode="complete", filter=flt)[1]["withheld"].values()) == {0}, flt
    _, p = build(w)
    # hand counts (C6): two raised flags (the gate's on beta-2025.pdf, the fixture's K1 flag), two archived documents
    # (the gate's acme-old.pdf, acme-archived.pdf), one deleted document (the gate's draft); the shareable CSV twin's
    # rows, private by P3 rule 3
    assert (p["withheld"]["flag"], p["withheld"]["archived"], p["withheld"]["deleted"]) == (2, 2, 1)
    assert p["withheld"]["same_bytes_private"] >= 1


# (c): the differential check (K12)

ID_KEYS = {
    **dict.fromkeys(("documents", "document_id", "evidence_document_id", "counter_document_id", "document"), "document"),
    **dict.fromkeys(("claim_id", "claim_a", "claim_b", "claims", "original_claim_id", "restated_claim_id"), "claim"),
    **dict.fromkeys(("version_id", "claim_version_id", "original_version_id", "restated_version_id", "low_version_id",
                     "high_version_id", "corrects_version"), "version"),
    **dict.fromkeys(("series_id", "series"), "series"),
    "restatement_id": "restatement", "adjudication_id": "adjudication", "expected_item_id": "expected item",
    "acceptance_id": "acceptance", "gap_confirmation_id": "gap confirmation", "pair_id": "pair",
}
ROW_IDS = {"documents": "document", "evidence_index": "claim", "related_claims": "claim", "restatement_log": "pair",
           "candidates": "pair", "evidence": "evidence", "disagreements": "version", "restatements": "restatement",
           "unpaired_restatements": "restatement", "adjudications": "adjudication", "expectations": "expected item",
           "declared": "document", "documents_checked": "document", "archived": "document", "deleted": "document",
           "accepted": "acceptance", "correction": "correction", "confirmation": "gap confirmation"}
ENTRY = re.compile(r"^(.+?) #(\d+) on version (\d+)(?: with claim (\d+) version (\d+))?$")


def renumber(d: dict) -> dict:
    """The dataset with every id (C3 id fields, and the ids GCT's own sentences carry) replaced by its rank of first
    appearance within its kind, and without data_version and built_at."""
    ranks = {}

    def n(kind, v):
        if v is None:
            return None
        r = ranks.setdefault(kind, {})
        return r.setdefault(v, len(r) + 1)

    def entry(s):
        m = ENTRY.match(s)
        if not m:
            return s
        kind, i, v, c, cv = m.groups()
        ns = kind if kind in ("correction", "adjudication") else ("record", kind)
        out = f"{kind} #{n(ns, int(i))} on version {n('version', int(v))}"
        return out + (f" with claim {n('claim', int(c))} version {n('version', int(cv))}" if c else "")

    def walk(x, table=None):
        if isinstance(x, list):
            return [walk(y, table) for y in x]
        if not isinstance(x, dict):
            return x
        if set(x) == {"id", "title", "rows", "parent_key", "columns"}:  # a sub-table
            return {**x, "rows": [walk(r, x["id"]) for r in x["rows"]]}
        out = {}
        for k, v in x.items():
            if k == "id" and isinstance(v, int):
                kind = ("record", x["kind"]) if table == "records" else ROW_IDS.get(table)
                out[k] = n(kind, v) if kind else v
            elif k in ID_KEYS and not (isinstance(v, (dict, str)) or isinstance(v, list) and v and isinstance(
                    v[0], dict)) or k in ("series_id", "series") and not isinstance(v, (dict, list)) or (
                    k == "series" and isinstance(v, list)):
                out[k] = [n(ID_KEYS[k], i) for i in v] if isinstance(v, list) else n(ID_KEYS[k], v)
            elif k == "evidence" and isinstance(v, list) and all(isinstance(i, int) for i in v):
                out[k] = [n("evidence", i) for i in v]
            elif k == "needs_renewed_review":
                out[k] = [entry(e) for e in v]
            elif k == "text" and table == "disagreements":
                out[k] = re.sub(r"on version (\d+)", lambda m: f"on version {n('version', int(m[1]))}", v)
            elif k == "evidence_text" and isinstance(v, str):
                out[k] = re.sub(r"mapping #(\d+)", lambda m: f"mapping #{n('mapping', int(m[1]))}", v)
            elif k in ("data_version", "built_at") and table is None:
                continue
            else:
                out[k] = walk(v, k if isinstance(v, (list, dict)) else table)
        return out

    return walk(d)


def diff(a, b, path="") -> list[str]:
    if type(a) is not type(b) or not isinstance(a, (dict, list)):
        return [] if a == b else [f"{path}: {a!r} != {b!r}"]
    if isinstance(a, dict):
        return ([f"{path}.{k}: on one side only" for k in sorted(set(a) ^ set(b), key=str)]
                + [x for k in a if k in b for x in diff(a[k], b[k], f"{path}.{k}")])
    if len(a) != len(b):
        return [f"{path}: {len(a)} items != {len(b)}"]
    return [x for i, (p, q) in enumerate(zip(a, b)) for x in diff(p, q, f"{path}[{i}]")]


@pytest.mark.parametrize("flt", FILTERS, ids=[json.dumps(x) for x in FILTERS])
def test_sanitised_w_equals_sanitised_w_prime(w, w2, env, flt):
    a, b = renumber(build(w, filter=flt)[0]), renumber(build(w2, filter=flt)[0])
    assert diff(a, b)[:40] == []


def test_the_differential_filters_depend_on_withheld_records(w, env):
    """B1: each review filter's and the unresolved filter's complete answer differs from its sanitised one (ruling 2:
    the fixture's nearest case for unresolved; the restatement filter's answer depends on no withheld record here)."""
    for flt in FILTERS[1:4]:
        c = {r["claim_id"] for r in build(w, mode="complete", filter=flt)[0]["evidence_index"] if r["label"] == "shareable"}
        s = {r["claim_id"] for r in build(w, filter=flt)[0]["evidence_index"]}
        assert c != s, flt


def test_the_third_company_and_its_canary_are_absent_from_both_sanitised_outputs(w, w2, env):
    """Ruling 3: W' keeps the third company (it names no document); neither sanitised output names it."""
    for root in (w, w2):
        d, _ = build(root)
        assert f.COMPANY not in text(d)
        assert not [x for x in d["documents"] if x["company"] == f.COMPANY]


def p_match_rows(d) -> dict:
    """{(document name, period label): (match state, match reason)} of acme-review.pdf and
    acme-printed.pdf's Scope 1 claims."""
    return {(r["document_name"], r["period_label"]): (r["match_state"], r["match_reason"]) for r in d["evidence_index"]
            if r["document_name"] in ("acme-review.pdf", "acme-printed.pdf") and r["metric_text"] == f.SCOPE_1}


def test_the_p_match_cells_are_recomputed_on_the_copy(w, w2, env):
    """Fixture cells A to D (worker-c-fixture.md): A reviewed into the consolidated series once the selected one is
    gone, B not linked once its only undecided candidate is gone, C into consolidated once the survivors reorder,
    D in the ordinary Works series by phase; each equal to W'."""
    s, s2 = (p_match_rows(build(r)[0]) for r in (w, w2))
    assert s == s2
    assert s["acme-review.pdf", f.BOTH[0]][0] == "reviewed" and s["acme-review.pdf", f.BOTH[1]] == (
        "not matchable", "not linked by review")
    assert s["acme-review.pdf", f.EARLIER[0]][0] == "reviewed" and s["acme-printed.pdf", f.EARLIER[0]][0] == "reviewed"
    c = p_match_rows(build(w, mode="complete")[0])
    assert c["acme-review.pdf", f.BOTH[1]][0] == "uncertain"  # W itself differs: the copy recomputed it


def test_the_reduced_copy_keeps_the_candidate_triggers_and_they_fire(w, env):
    """Moka: projection 16's three triggers survive deserialize on the copy, and dropping a withheld claim's rows
    there removes its candidate rows (and a removed series' rows)."""
    from gct.reports import pin
    from gct.stages import view

    conn = store.connect(w / layout.DB_NAME)
    pconn, ident = pin.pin(conn, w)
    rconn = dataset._Build(conn, w, pconn, ident, view.normalise({}), []).reduce()
    try:
        names = {n for (n,) in rconn.execute("SELECT name FROM sqlite_master WHERE type = 'trigger'")}
        assert {"match_candidate_claim", "match_candidate_series", "match_candidate_rule"} <= names
        claims = {c for (c,) in rconn.execute("SELECT id FROM claim")}
        series = {s for (s,) in rconn.execute("SELECT id FROM series")}
        before = pconn.execute("SELECT claim_id, series_id FROM match_candidate").fetchall()
        gone = [r for r in before if r["claim_id"] not in claims or r["series_id"] not in series]
        assert gone  # the fixture's candidates of removed series (cells A to C); no waiting claim is withheld itself
        after = {tuple(r) for r in rconn.execute("SELECT claim_id, series_id FROM match_candidate")}
        assert after == {tuple(r) for r in before} - {tuple(r) for r in gone}
        # The claim trigger on a copy made the same way (C5a step 1), since the fixture withholds no waiting claim.
        copy = sqlite3.connect(":memory:")
        copy.deserialize(pconn.serialize())
        (cid,) = copy.execute("SELECT claim_id FROM match_candidate LIMIT 1").fetchone()
        copy.execute("DELETE FROM claim WHERE id = ?", (cid,))
        assert copy.execute("SELECT COUNT(*) FROM match_candidate WHERE claim_id = ?", (cid,)).fetchone()[0] == 0
        assert copy.execute("SELECT COUNT(*) FROM match_candidate").fetchone()[0] > 0
        copy.close()
    finally:
        for c in (rconn, pconn, conn):
            c.close()


def test_restatement_rows_attach_through_their_claims_in_both_modes(w, w2, env):
    """Scenario 2 (an earlier version, still paired) and the two unpaired rows (one citing a private document)."""
    for mode in dataset.MODES:
        d, _ = build(w, mode=mode)
        earlier = [r for p in d["restatement_log"] for r in p["restatements"]["rows"] if not r["binds"]]
        assert earlier and all(r["earlier"] == dataset.EARLIER for r in earlier), mode
        unpaired = d["unpaired_restatements"]["rows"]
        assert len(unpaired) == (2 if mode == "complete" else 1), mode
    a, b = (renumber(build(r)[0])["unpaired_restatements"] for r in (w, w2))
    assert a == b


# Part B (story 7.7): every output and both packages of tests/leakage_fixture.py, scanned from the files alone

FORMATS = ("json", "xlsx", "csv", "pdf", "page")
PACKAGES = ("complete package", "sanitised package")
TABLES = ["evidence_index", "restatement_log", "gap_report"]  # what gct report all renders
GLOBS = {"json": ("dataset.json", "*.schema.json"), "xlsx": ("*.xlsx",), "csv": ("*/*.csv",), "pdf": ("*.pdf",),
         "page": ("*.html", "crops/*.png")}


def rgb(png: bytes) -> str:
    from PIL import Image

    return hashlib.sha256(Image.open(io.BytesIO(png)).convert("RGB").tobytes()).hexdigest()


def with_crops(d, c) -> list:
    """leakage_fixture.SEEDS plus one image seed per private page entry's crop (7.1 G: the crop hashes are computed at
    test time from crops.cut); a crop whose pixels equal a seed's value (IMG-PAGE-01) is kept once."""
    seeds = copy.deepcopy(lf.SEEDS)
    private = {x["id"] for x in d["documents"] if x["label"] == "private"}
    for e in d["evidence"]:
        if e["page"] is not None and e["document_id"] in private and e["id"] in c and rgb(c[e["id"]]) not in {
                s["value"] for s in seeds}:
            seeds.append({"name": f"LEAK-CROP-E{e['id']:04d}", "kind": "image", "expect": "complete_only",
                          "value": rgb(c[e["id"]])})
    return seeds


def files_of(folder, fmt) -> list[Path]:
    return sorted({p for g in GLOBS[fmt] for p in Path(folder).glob(g) if p.is_file()})


def scan(paths, seeds) -> dict:
    """found over several files: {name: ["<file>: <place>"]}."""
    out = {}
    for p in paths:
        for name, places in leakage.found(p, seeds).items():
            out.setdefault(name, []).extend(f"{p.name}: {x}" for x in places)
    return out


def texts(paths) -> str:
    """Every text part of the files, one per line, white space squeezed out (a PDF may break a long token)."""
    return "\n".join("".join(v.split()) for p in paths for k, _, v in leakage.parts(p) if k == "text")


def has(paths, literal: str) -> bool:
    return "".join(literal.split()) in texts(paths)


def say(*argv) -> str:
    """One command's terminal text (no --json)."""
    return answered(*argv)[0]


def answered(*argv) -> tuple[str, str]:
    """(stdout, stderr) of one command that must succeed; stderr caught both as sys.stderr and as file descriptor 2,
    so a print to either is kept (Codex fix check 1, R2-1)."""
    out, err = io.StringIO(), io.StringIO()
    with tempfile.TemporaryFile() as fd2:
        saved = os.dup(2)
        os.dup2(fd2.fileno(), 2)
        try:
            with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
                code = cli.main([str(a) for a in argv])
        finally:
            os.dup2(saved, 2)
            os.close(saved)
        fd2.seek(0)
        errors = err.getvalue() + fd2.read().decode("utf-8", "replace")
    assert code == 0, (argv, out.getvalue(), errors)
    return out.getvalue(), errors


def made(root) -> dict:
    """Every report format in both modes through gct report all, and both packages through gct package create (the
    complete one with --include-private): {mode: export folder, package column: package folder, "responses": [(the
    command, its --json answer, its stderr)] (renderer notes, report lines, package preview and creation lines)}."""
    at, out, responses = ("--workspace", root), {}, []

    def run(*argv):
        text, errors = answered(*argv, *at, "--json")
        answer = json.loads(text)
        responses.append((" ".join(argv[:2]) + (" --complete" if "--complete" in argv else ""), answer, errors))
        return answer

    for mode, flag in (("sanitised", ()), ("complete", ("--complete",))):
        out[mode] = Path(run("report", "all", "--format", ",".join(FORMATS), *flag)["folder"])
    token = run("package", "preview")["token"]
    out["sanitised package"] = root / run("package", "create", "--confirm", token)["folder"]
    token = run("package", "preview", "--complete")["token"]
    out["complete package"] = root / run("package", "create", "--complete", "--include-private", "--confirm",
                                         token)["folder"]
    return {**out, "responses": responses}


@pytest.fixture(scope="module")
def lwm(env, tmp_path_factory):
    """W of 7.7: the workspace scanned before any output exists (AC 8's source proof), then made(); with the complete
    dataset as exported, its crops and the seeds."""
    root = lf.build(tmp_path_factory.mktemp("lw"))
    gc.collect()  # closes in-process SQLite connections: their -shm and -wal files settle before found lists them
    out = {"root": root, "source": leakage.found(root, lf.SEEDS), **made(root)}
    d = json.loads((out["complete"] / "dataset.json").read_text(encoding="utf-8"))
    c = crops.cut(root, d)
    return {**out, "dataset": d, "crops": c, "seeds": with_crops(d, c)}


@pytest.fixture(scope="module")
def lw(lwm):
    """lwm and every output's found result ("hits")."""
    hits = {(fmt, mode): scan(files_of(lwm[mode], fmt), lwm["seeds"]) for fmt in FORMATS for mode in dataset.MODES}
    hits.update({p: leakage.found(lwm[p], lwm["seeds"]) for p in PACKAGES})
    return {**lwm, "hits": hits}


@pytest.fixture(scope="module")
def lw2(env, tmp_path_factory):
    return {"root": lf.build(tmp_path_factory.mktemp("lw2"), without_private=True)}


# AC 7, 8: the seed matrix (section 5), as delivered. A cell is C (found in the complete output, absent from the
# sanitised one), N (absent from both), n/c with its reason (absent from both: a format that starts carrying the seed
# fails, and the matrix is updated on purpose); a package cell is found or absent, 7.6's rule cited where it decides.
NO_IMAGE = "n/c: the format holds no images"
NO_EVIDENCE = "n/c: cell content and sha256 sit in dataset.evidence, which 7.3 does not print"
NO_SHA = "n/c: an evidence entry is printed by its document, version and page, never its sha256"
IMAGE = {**dict.fromkeys(("json", "xlsx", "csv"), NO_IMAGE),
         "complete package": "absent: 7.6 ships no page images, and the original draws the marker as a vector"
                             " rectangle, which found does not render"}
CELLS = {
    "CON-ACCEPT-01": dict.fromkeys(FORMATS, "n/c: C3c accepted carries acceptance ids only (K7)"),
    "DOC-BASIS-01": dict.fromkeys(FORMATS, "n/c: no dataset field carries a basis wording (7.1 G (a))"),
    "FILE-PDF-02": dict.fromkeys(FORMATS, "n/c: an archived document has no evidence entry, so no sha256"),
    **{n: {"xlsx": NO_EVIDENCE, "csv": NO_EVIDENCE, "pdf": NO_SHA, "page": NO_SHA}
       for n in ("FILE-PDF-01", "FILE-CSV-01", "FILE-PDF-03", "FILE-PDF-04")},
    **{n: {"xlsx": NO_EVIDENCE, "csv": NO_EVIDENCE} for n in ("DOC-CELL-02", "LEAK-11")},  # K5, and L13's heading
    "LEAK-31": {"complete package": "found: a complete package ships a document's source (P5)"},
    "LEAK-38": {"complete package": "found: a complete package ships contributor identity"},
    **{n: {"complete package": "7.6 found: a complete package may hold the run log's path lines (7.6 D6)"}
       for n in ("LEAK-37", "PATH-01")},
    "LEAK-35": {"complete package": "7.6 absent: a package holds no logs/ folder"},
    "LEAK-36": {"complete package": "7.6 absent: a package holds no files/.incoming-* file"},
}


def cell(seed, column) -> str:
    over = {**(IMAGE if seed["kind"] == "image" else {}), **CELLS.get(seed["name"], {})}
    if column in over:
        return over[column]
    if column in PACKAGES:
        return "found" if column == "complete package" and seed["expect"] == "complete_only" else "absent"
    return "C" if seed["expect"] == "complete_only" else "N"


def test_every_stored_never_seed_is_in_the_source_workspace(lw):
    """AC 8: each never seed of leakage_fixture.STORED_RAW is found in the workspace before any output is built, so
    its absence from the outputs is meaningful; CRED-RUN-01 (scrubbed when written) and CRED-ENV-01 (environment only)
    are never stored. The scan raised no Unreadable (found raises on any part it cannot read)."""
    assert lf.STORED_RAW <= set(lw["source"]), sorted(lf.STORED_RAW - set(lw["source"]))
    assert not {"CRED-RUN-01", "CRED-ENV-01"} & set(lw["source"])


@pytest.mark.parametrize("column", FORMATS + PACKAGES)
def test_every_cell_of_the_seed_matrix(lw, column):
    """AC 7, 8 (and 10's marker cell): every seed, crop seeds included, against its cell; a C cell is a find in the
    complete output, so the scan could have failed."""
    hits, wrong = lw["hits"], []
    for s in lw["seeds"]:
        want = cell(s, column)
        if column in PACKAGES:
            got = "found" if s["name"] in hits[column] else "absent"
            if want.split(":")[0].split()[-1] != got:
                wrong.append((s["name"], want, hits[column].get(s["name"])))
        else:
            got = (s["name"] in hits[column, "complete"], s["name"] in hits[column, "sanitised"])
            if got != ((True, False) if want == "C" else (False, False)):
                wrong.append((s["name"], want, hits[column, "complete"].get(s["name"]),
                              hits[column, "sanitised"].get(s["name"])))
    assert wrong == []


@pytest.mark.parametrize("fmt", ("pdf", "page"))
def test_each_crop_seed_and_the_marker_only_where_that_crop_is_drawn(lw, fmt):
    """AC 6, 7: in each complete PDF and in the complete page folder, the image seeds found are exactly those of the
    crops it must draw and the marker seed only where the private page's crop or whole page is drawn (7.4's and
    7.5's expected_hits, read only)."""
    d, c, seeds = lw["dataset"], lw["crops"], lw["seeds"]
    images = {s["name"] for s in seeds if s["kind"] == "image"}
    if fmt == "pdf":
        import test_render_pdf as m

        wants = {m.FILES[r]: m.expected_hits(d, c, r, seeds) for r in m.REPORTS}
    else:
        import test_render_page as m

        wants = {None: m.expected_hits(d, c, seeds)}
    seen = set()
    for name, want in wants.items():
        paths = [lw["complete"] / name] if name else files_of(lw["complete"], fmt)
        got = set(scan(paths, seeds)) & images
        assert got == {n for n in want if n in images}, (name, sorted(got ^ {n for n in want if n in images}))
        seen |= got
    assert "IMG-MARKER-01" in seen and len(seen) > 2  # the check can fail: crops and the marker are drawn


def test_the_pages_in_a_browser_show_and_request_nothing_private_or_outside(lw, tmp_path):
    """AC 11: the complete and sanitised pages opened from file:// in headless Chrome (7.5's
    tests/report_page_browser.mjs): the complete innerText holds exactly the text seeds the pages print (7.5's
    expected_hits), the sanitised one none, and no page requests a URL outside its own folder (K22)."""
    import test_render_page as m

    urls = {mode: [(lw[mode] / m.FILES[r]).as_uri() for r in m.REPORTS] for mode in dataset.MODES}
    out = m.browse(tmp_path, *urls["complete"], *urls["sanitised"])
    assert [p["url"] for p in out["pages"]] == urls["complete"] + urls["sanitised"]
    assert out["requests"] and all(any(u.startswith(lw[mode].as_uri() + "/") for mode in dataset.MODES)
                                   for u in out["requests"]), out["requests"]
    for mode in dataset.MODES:
        n = len(m.REPORTS)
        pages = out["pages"][:n] if mode == "complete" else out["pages"][n:]
        (tmp_path / f"{mode}.txt").write_text("\n".join(p["text"] for p in pages), encoding="utf-8")
    want = {k for k, how in m.expected_hits(lw["dataset"], lw["crops"], lw["seeds"]).items() if how == "text"}
    assert set(leakage.found(tmp_path / "complete.txt", lw["seeds"])) == want and want
    assert leakage.found(tmp_path / "sanitised.txt", lw["seeds"]) == {}


def test_the_positive_literals_ship_where_the_matrix_says(lw):
    """L14's shareable neighbour cell ships in both modes wherever a format prints cell content (n/c in XLSX and CSV,
    as K5) and in both packages; the pseudonym is in every sanitised output and in no complete one. Neither is passed
    to found."""
    for fmt in FORMATS:
        both = (has(files_of(lw["complete"], fmt), lf.NEIGHBOUR), has(files_of(lw["sanitised"], fmt), lf.NEIGHBOUR))
        assert both == ((False, False) if fmt in ("xlsx", "csv") else (True, True)), fmt
        named = tuple(has(files_of(lw[m], fmt), lf.PSEUDONYM) for m in ("complete", "sanitised"))
        assert named == (False, True), fmt
    for p in PACKAGES:  # the shareable workbook's original
        assert any(lf.NEIGHBOUR.encode() in x.read_bytes() for x in (lw[p] / "files").rglob("*") if x.is_file()), p


def test_the_preview_texts_are_in_no_output(lw):
    """AC 8: the preview sentence and the relabel command are in no output and no package, and they are in the
    gct report preview text of a selection of private documents only (they never sit in the workspace)."""
    private = next(x["id"] for x in lw["dataset"]["documents"] if x["name"] == f.NAME)
    shown = say("report", "preview", "--filter", json.dumps({"document": [private]}), "--workspace", lw["root"])
    for line in (dataset.ALL_PRIVATE, dataset.RELABEL):
        assert line in shown
        for fmt in FORMATS:
            for mode in dataset.MODES:
                assert not has(files_of(lw[mode], fmt), line), (fmt, mode, line)
        for p in PACKAGES:
            assert not has([lw[p]], line), (p, line)


def local_paths(base) -> list:
    return [re.compile(r"/(?:Users|home)/[A-Za-z0-9._-]+/"),  # anyone_home, test_hygiene.py:88
            re.compile(re.escape("".join(str(base).split()))),
            re.compile(re.escape("".join(os.path.realpath(tempfile.gettempdir()).split())))]


def test_no_local_path_in_any_report_output_or_the_sanitised_package(lw, tmp_path_factory):
    """Section 3's non-seed path checks: anyone_home, the test's temporary root and the realpath of the temporary
    folder in no text part of any report output (both modes) or the sanitised package; each is first found in the
    source workspace's logs and run lines (the L8 error log's traceback names GCT's own files). The complete
    package may hold the run log's path lines (7.6 D6)."""
    source = texts([lw["root"] / "logs", lw["root"] / layout.DB_NAME])
    patterns = local_paths(tmp_path_factory.getbasetemp())
    assert all(p.search(source) for p in patterns), [p.pattern for p in patterns if not p.search(source)]
    for fmt in FORMATS:
        for mode in dataset.MODES:
            t = texts(files_of(lw[mode], fmt))
            assert not [p.pattern for p in patterns if p.search(t)], (fmt, mode)
    t = texts([lw["sanitised package"]])
    assert not [p.pattern for p in patterns if p.search(t)]


WROTE = re.compile(r"Wrote (?:sanitised|complete) report data for data version \d+ to (.+):")  # cli.cmd_report


def test_no_credential_or_path_seed_in_any_command_response(lw, tmp_path, tmp_path_factory):
    """Codex review F1 and fix check 1 (R2-1, R2-2): the answers the report and package commands give, as --json
    (made()'s) and as terminal text (the same commands on a copy of W), in both modes. A report names the folder it
    wrote in exactly two places, the --json "folder" field and the "Wrote ... to FOLDER:" heading line; those two are
    checked to be that folder and taken out by field. Everything else (renderer notes, report lines, package preview
    and creation lines) is scanned as it is: no never seed of kind credential or path, no local path (anyone_home,
    the workspace, the temporary folder) and, in a sanitised command's answer, no original's file hash as text (they
    are complete_only seeds: a complete answer names them); the local preview may name private
    documents, so other seeds are not looked for there (the matrix covers the files). Every command's stderr, caught
    apart, holds no never seed at all."""
    never = [s for s in lw["seeds"] if s["expect"] == "never"]
    seeds = [s for s in never if s["kind"] in ("credential", "path")]
    assert {"credential", "path"} == {s["kind"] for s in seeds}
    hashes = [s["value"] for s in lw["seeds"] if s["kind"] == "file_hash"]
    copy = Path(shutil.copytree(lw["root"], tmp_path / "ws"))
    at, text = ("--workspace", copy), []
    for flag in ((), ("--complete",)):
        text.append((f"report all{' '.join(('', *flag))}",
                     *answered("report", "all", "--format", ",".join(FORMATS), *flag, *at)))
    text.append(("package preview", *answered("package", "preview", *at)))
    text.append(("package preview --complete", *answered("package", "preview", "--complete", *at)))
    text.append(("package create", *answered("package", "create", "--confirm",
                                             gct("package", "preview", *at)["token"], *at)))
    token = gct("package", "preview", "--complete", *at)["token"]
    text.append(("package create --complete",
                 *answered("package", "create", "--complete", "--include-private", "--confirm", token, *at)))
    answers, errors = tmp_path / "answers", tmp_path / "stderr"
    answers.mkdir()
    errors.mkdir()
    for n, (root, (command, answer, stderr)) in enumerate([*((lw["root"], x) for x in lw["responses"]),
                                                           *((copy, x) for x in text)]):
        def written(folder):  # the folder a report wrote: under its own workspace's exports/
            return Path(os.path.realpath(folder)).parent == Path(os.path.realpath(root)) / "exports"

        if isinstance(answer, dict):  # --json
            answer = dict(answer)
            assert not command.startswith("report") or written(answer.pop("folder")), command
            answer = json.dumps(answer, ensure_ascii=False)
        elif command.startswith("report"):
            heading = [m for line in answer.splitlines() if (m := WROTE.fullmatch(line))]
            assert len(heading) == 1 and written(heading[0][1]), command
            answer = answer.replace(heading[0][0], "")
        (answers / f"{n:02d} {command}.txt").write_text(answer, encoding="utf-8")
        (errors / f"{n:02d} {command}.txt").write_text(stderr, encoding="utf-8")
        local = local_paths(tmp_path_factory.getbasetemp()) + [re.compile(re.escape(str(root))),
                                                        re.compile(re.escape(os.path.realpath(root)))]
        assert not [p.pattern for p in local if p.search(answer)], command
        assert "--complete" in command or not [h for h in hashes if h in answer], command  # complete_only seeds
    assert len(list(answers.iterdir())) == 12 == len(list(errors.iterdir()))
    assert leakage.found(answers, seeds) == {}
    assert leakage.found(errors, never) == {}


MARKS = ("history_mark", "labels_mark", "hidden_by_archive", "History mark", "Labels mark")


def test_the_history_and_labels_marks_are_in_complete_report_outputs_only(lw):
    """L18 (K2, K10, K13): the complete dataset's dotted history_mark and the identity labels are in every complete
    report output, and neither they nor the key names are in a sanitised one; in JSON the key names and labels are
    looked for outside words and columns (ruling 4), in printed formats everywhere."""
    mark = lw["dataset"]["data_version"]["history_mark"]
    assert mark.count(".") > 5
    for fmt in FORMATS:
        if fmt == "json":
            d = json.loads((lw["sanitised"] / "dataset.json").read_text(encoding="utf-8"))
            t = json.dumps({k: v for k, v in d.items() if k not in ("words", "columns")}, ensure_ascii=False)
        else:
            t = texts(files_of(lw["sanitised"], fmt))
        assert not [x for x in (mark, *MARKS) if "".join(x.split()) in "".join(t.split())], fmt
        assert has(files_of(lw["complete"], fmt), mark), fmt
        assert fmt == "json" or all(has(files_of(lw["complete"], fmt), x) for x in MARKS[3:]), fmt


@pytest.mark.parametrize("fmt", FORMATS)
def test_a_private_row_put_back_into_a_sanitised_dataset_is_found(lw, fmt, tmp_path):
    """AC 9: a builder bug simulated: one private evidence-index row, its evidence entries and its document's entry put
    back into the sanitised dataset, rendered with the format's write; its seeds are found, and for the crop formats
    (crops cut from the injected dataset) the private crop's pixel hash and the marker too."""
    s, _ = build(lw["root"])
    c = lw["dataset"]
    row = next(r for r in c["evidence_index"] if r["metric_text"].startswith(f.METRIC))
    entries = {e["id"]: e for e in c["evidence"]}
    new = {i: max((e["id"] for e in s["evidence"]), default=0) + n for n, i in enumerate(row["evidence"], 1)}
    s["evidence"] += [{**entries[i], "id": new[i]} for i in row["evidence"]]
    s["evidence_index"].append({**copy.deepcopy(row), "evidence": [new[i] for i in row["evidence"]]})
    s["documents"].append(next(x for x in c["documents"] if x["id"] == row["document_id"]))
    k = crops.cut(lw["root"], s) if export.FORMATS[fmt]["crops"] else {}
    export.FORMATS[fmt]["write"](s, TABLES, tmp_path, k)
    got = leakage.found(tmp_path, lw["seeds"])
    want = {"DOC-METRIC-01", "DOC-NAME-01"}
    if k:
        want |= {"IMG-MARKER-01"} | {x["name"] for x in lw["seeds"] if x["kind"] == "image"
                                     and x["value"] in {rgb(k[i]) for i in new.values() if i in k}}
        assert len(want) > 3
    assert want <= set(got), sorted(want - set(got))


def test_the_packages_hold_no_wal_shm_incoming_or_exports_file_and_no_private_original_ships_sanitised(lw):
    """AC 10: neither package holds a -wal, -shm, .incoming-* or exports/ file; no file of the sanitised package has
    the sha256 of a private original (the L2 twin shares FILE-PDF-01's bytes); the complete package holds them, made
    with 7.6's authorisation (--include-private). The marker cells are in the matrix. Cited, not repeated:
    tests/test_package.py::test_a_sanitised_package_holds_what_its_preview_says_and_reopens,
    ::test_a_complete_package_needs_include_private_and_carries_every_row, ::test_only_package_py_holds_the_phrase and
    ::test_no_command_calls_a_copied_or_changed_folder_sanitised (FR72)."""
    private = {x["sha256"] for x in gct("doc", "list", "--workspace", lw["root"])["documents"]
               if x["privacy"] == "private" and x["sha256"]}
    held = {}
    for p in PACKAGES:
        names = [x.relative_to(lw[p]).as_posix() for x in lw[p].rglob("*") if x.is_file()]
        assert not [n for n in names if n.endswith(("-wal", "-shm")) or ".incoming-" in n or "exports/" in n], p
        held[p] = {hashlib.sha256((lw[p] / n).read_bytes()).hexdigest() for n in names}
    assert not held["sanitised package"] & private
    assert next(s["value"] for s in f.SEEDS if s["name"] == "FILE-PDF-01") in held["complete package"] & private


# AC 7a: the differential check (K12) on every rendered format, beside the seed scan. Clock parts, named up front:
# an XLSX file's raw bytes and file hash (every zip member's date_time) and the parts of its docProps/core.xml
# (openpyxl stamps dcterms:created and dcterms:modified with the current time). Those parts are left out, and in
# their place compared() compares what they hold besides the clock (Codex review F4): core.xml's bytes with only
# the values of dcterms:created and dcterms:modified replaced, and each XLSX's zip fields but date_time.
CLOCK = re.compile(r"^[^>]*\.xlsx(?:: raw bytes| > docProps/core\.xml(?:: (?:raw bytes|text|XML|HTML text))?)?$")
STAMPS = re.compile(rb"(<dcterms:(created|modified)\b[^>]*>)[^<]*(</dcterms:\2>)")


def normalised(root, mode="sanitised") -> dict:
    """renumber's dataset, keeping data_version (sha256 a placeholder) and built_at, which the renderers print."""
    d, _ = build(root, mode=mode)
    return {**renumber(d), "data_version": {**d["data_version"], "sha256": "0" * 64}, "built_at": d["built_at"]}


def render(root, d, fmt, name) -> Path:
    """write into a folder inside the workspace, crops cut from that workspace's own originals."""
    folder = Path(root) / "k12" / name / fmt
    folder.mkdir(parents=True)
    export.FORMATS[fmt]["write"](d, TABLES, folder, crops.cut(root, d) if export.FORMATS[fmt]["crops"] else {})
    return folder


def compared(folder) -> list:
    """The parts of a folder's files to compare: every part but the CLOCK ones, and for each XLSX its core.xml with
    the two stamps' values replaced and its zip fields but date_time (archive comment; per member name, comment,
    extra field, compression, attributes and size)."""
    out = [p for p in leakage.parts(folder) if not CLOCK.match(p[1])]
    for x in sorted(Path(folder).rglob("*.xlsx")):
        name = x.relative_to(folder).as_posix()
        with zipfile.ZipFile(x) as z:
            core = STAMPS.sub(rb"\1<stamp>\3", z.read("docProps/core.xml"))
            fields = [z.comment] + [(i.filename, i.comment, i.extra, i.compress_type, i.external_attr, i.file_size)
                                    for i in z.infolist()]
        out += [("text", f"{name} > docProps/core.xml: stamps replaced", core.decode("utf-8")),
                ("text", f"{name}: zip fields but date_time", repr(fields))]
    return sorted(out)


@pytest.fixture(scope="module")
def k12(lw, lw2):
    """{(mode, workspace): normalised dataset} of W and W', with their version numbers equal."""
    out = {(m, x): normalised(ws["root"], m) for m in dataset.MODES for x, ws in (("W", lw), ("W'", lw2))}
    assert out["sanitised", "W"]["data_version"]["version"] == out["sanitised", "W'"]["data_version"]["version"]
    return out


@pytest.mark.parametrize("fmt", FORMATS)
def test_sanitised_w_and_w_prime_render_equal_parts_and_complete_ones_differ(lw, lw2, k12, fmt):
    """AC 7a: each format rendered from the normalised sanitised(W) and sanitised(W') has equal text parts and image
    pixel hashes, the named clock parts excluded (and seed-scanned); the same comparison of complete(W) and
    complete(W') differs (the check can fail). The K1 claim is in the comparison with no exemption."""
    a, b = (render(ws["root"], k12["sanitised", x], fmt, "sanitised") for x, ws in (("W", lw), ("W'", lw2)))
    pa, pb = compared(a), compared(b)
    assert pa == pb, (sorted(set(pa) - set(pb))[:5], sorted(set(pb) - set(pa))[:5])
    assert leakage.found(a, lw["seeds"]) == {} == leakage.found(b, lw["seeds"])
    ca, cb = (render(ws["root"], k12["complete", x], fmt, "complete") for x, ws in (("W", lw), ("W'", lw2)))
    assert compared(ca) != compared(cb)


def test_non_clock_xlsx_metadata_that_differs_between_w_and_w_prime_fails_the_differential(lw, lw2, k12):
    """Codex review F4's mutation, in the test: the sanitised XLSX of W and of W' compare equal; then W's files get
    a core.xml description (the reviewer's "Private documents: N"), a changed dcterms:creator, or a zip comment,
    one at a time, and each makes compared() differ, though no seed is in them."""
    a, b = (render(ws["root"], k12["sanitised", x], "xlsx", "metadata") for x, ws in (("W", lw), ("W'", lw2)))
    assert compared(a) == compared(b)
    description = b'<dc:description xmlns:dc="http://purl.org/dc/elements/1.1/">Private documents: 16</dc:description>'
    for change in (lambda n, d: (d.replace(b"</cp:coreProperties>", description + b"</cp:coreProperties>")
                                 if n == "docProps/core.xml" else d),
                   lambda n, d: d.replace(b"<dc:creator>openpyxl", b"<dc:creator>Acme") if n == "docProps/core.xml"
                   else d, "comment"):
        for x in sorted(a.glob("*.xlsx")):
            data, buf = x.read_bytes(), io.BytesIO()
            with zipfile.ZipFile(io.BytesIO(data)) as src, zipfile.ZipFile(buf, "w") as dst:
                for i in src.infolist():
                    dst.writestr(i, src.read(i) if change == "comment" else change(i.filename, src.read(i)))
                dst.comment = b"Private documents: 16" if change == "comment" else src.comment
            x.write_bytes(buf.getvalue())
            assert compared(a) != compared(b), change
            assert leakage.found(a, lw["seeds"]) == {}
            x.write_bytes(data)
            break


def test_the_named_clock_parts_differ_when_the_clock_moves_and_no_other_part_does(lw, k12, tmp_path, monkeypatch):
    """AC 7a: two renders of one normalised dataset, the clock moved 400 days between them (time.time for zipfile's
    member date_time, openpyxl.packaging.core's datetime for core.xml): every XLSX part CLOCK names differs and
    nothing else does, and compared() of the two renders is equal (only dcterms:created, dcterms:modified and
    date_time moved); every zip member's date_time differs; JSON, CSV, PDF and page have no clock part."""
    d = k12["sanitised", "W"]

    def write(name):
        for fmt in FORMATS:
            (tmp_path / name / fmt).mkdir(parents=True)
            export.FORMATS[fmt]["write"](d, TABLES, tmp_path / name / fmt,
                                         crops.cut(lw["root"], d) if export.FORMATS[fmt]["crops"] else {})
        return {fmt: leakage.parts(tmp_path / name / fmt) for fmt in FORMATS}

    first = write("first")
    later = time.time() + 400 * 86400

    class Later(datetime.datetime):
        @classmethod
        def now(cls, tz=None):
            return datetime.datetime.fromtimestamp(later, tz)

    import openpyxl.packaging.core as core

    monkeypatch.setattr(core, "datetime", types.SimpleNamespace(datetime=Later, timezone=datetime.timezone))
    monkeypatch.setattr(time, "time", lambda: later)
    second = write("second")
    monkeypatch.undo()
    for fmt in FORMATS:
        moved = {(k, p) for k, p, _ in set(first[fmt]) ^ set(second[fmt])}
        clock = {(k, p) for k, p, _ in first[fmt] if CLOCK.match(p)}
        assert moved == clock, (fmt, sorted(moved ^ clock)[:5])
        assert bool(clock) == (fmt == "xlsx"), fmt
        # what compared() keeps of the clock parts (core.xml but its two stamps, the zip fields) did not move
        a, b = compared(tmp_path / "first" / fmt), compared(tmp_path / "second" / fmt)
        assert a == b, (fmt, sorted(set(a) ^ set(b))[:3])
    for x in files_of(tmp_path / "first" / "xlsx", "xlsx"):
        with zipfile.ZipFile(x) as a, zipfile.ZipFile(tmp_path / "second" / "xlsx" / x.name) as b:
            assert all(i.date_time != j.date_time for i, j in zip(a.infolist(), b.infolist())), x.name
        names = {p for _, p, _ in first["xlsx"] if CLOCK.match(p) and p.startswith(x.name)}
        assert {f"{x.name}", f"{x.name}: raw bytes", f"{x.name} > docProps/core.xml"} <= names, names


# AC 12: behaviour (section 4) across the five report outputs; packages by 7.6's tests and the matrix

def module(fmt):
    """The format's own test module, which holds its read_back and printed (imported late: each imports this one)."""
    import importlib

    return importlib.import_module({"json": "test_export_json", "xlsx": "test_render_sheets",
                                    "csv": "test_render_sheets", "pdf": "test_render_pdf",
                                    "page": "test_render_page"}[fmt])


def back(fmt, folder) -> dict:
    """read_back of an export folder through the format's module: {"data_version": [identity of each file read],
    "rows": {table: rows}}; XLSX, CSV and PDF are read file by file, JSON and the page folder whole."""
    m = module(fmt)
    if fmt in ("json", "page"):
        b = m.read_back(folder)
        return {"data_version": [b["data_version"]], "rows": b["rows"]}
    out = {"data_version": [], "rows": {}}
    for r in TABLES:
        stem = r.replace("_", "-")
        b = m.read_back(folder / {"xlsx": f"{stem}.xlsx", "csv": stem, "pdf": f"{stem}.pdf"}[fmt])
        out["data_version"].append(b["data_version"])
        out["rows"].update(b["rows"])
    return out


def exported(lw, mode) -> dict:
    return json.loads((lw[mode] / "dataset.json").read_text(encoding="utf-8"))


def keys_of(x) -> set:
    if isinstance(x, dict):
        return set(x) | {k for v in x.values() for k in keys_of(v)}
    return {k for v in x for k in keys_of(v)} if isinstance(x, list) else set()


COUNTS = ("hidden_by_archive", "deleted_in_filter", "matched", "total", "excluded", "history_mark", "labels_mark")


def test_archive_counts_and_notices_in_every_report_output(lw):
    """Archive and counts (K10, K2, P8): every notice of each mode's dataset prints in that mode's output of every
    format, the sanitised ones with the privacy and archive notices; the sanitised dataset carries none of the count
    or mark keys outside words and columns; a review on a shareable claim citing the archived private document is
    in the complete dataset with its evidence archived and in no sanitised one (the archived name is the matrix's
    DOC-NAME-02 row)."""
    for mode in dataset.MODES:
        d = exported(lw, mode)
        for fmt in FORMATS:
            assert [n for n in d["notices"] if not has(files_of(lw[mode], fmt), n)] == [], (mode, fmt)
        records = [r for x in d["evidence_index"] for r in x["records"]["rows"]]
        assert any(r["evidence_state"] == "archived" for r in records) == (mode == "complete")
    s = exported(lw, "sanitised")
    assert {dataset.PRIVACY_NOTICE, dataset.ARCHIVE_NOTICE} <= set(s["notices"])
    assert not keys_of({k: v for k, v in s.items() if k not in ("words", "columns")}) & set(COUNTS)
    assert {"history_mark", "labels_mark"} <= keys_of(exported(lw, "complete")["data_version"])


def test_a_deleted_documents_seeds_are_in_the_workspace_before_the_deletion(env, tmp_path):
    """Permanent deletion (L6): with the deletion left out, L6's seeds are found in the workspace, so their N cells
    in the matrix (every output and package built after the deletion) can fail. The L7 file in exports/ is in no
    package (the packages test)."""
    root = lf.build(tmp_path, delete=False)
    gc.collect()  # as in lw
    assert {"LEAK-32", "LEAK-33"} <= set(leakage.found(root, lf.SEEDS))


def test_an_unfinished_deletion_refuses_every_report_and_writes_nothing_until_it_is_finished(lw, tmp_path,
                                                                                          monkeypatch):
    """Archived, then deleted with its purge pending (K3): gct report refuses every format with "Document N was
    deleted but its removal did not finish." and exports/ is unchanged (names and hashes); after the finishing
    confirm the next report is written and its sanitised outputs hold no seed."""
    from gct.stages import deletion

    root = Path(shutil.copytree(lw["root"], tmp_path / "ws"))
    at, who = ("--workspace", root), ("--as", "kim", "--capacity", f.CAPACITY, "--rationale", "the client asked")
    doc = next(x["id"] for x in gct("doc", "list", *at)["documents"] if x["state"] == "archived"
               and x["privacy"] == "private")

    def boom(*a):
        raise OSError("held open")

    monkeypatch.setitem(deletion.WORK, "published", boom)
    token = gct("doc", "delete", doc, *at, ok=4)["token"]
    gct("doc", "delete", doc, "--confirm", token, *who, *at, ok=4)
    monkeypatch.undo()

    def listing():
        return {p.relative_to(root).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest()
                for p in (root / "exports").rglob("*") if p.is_file()}

    before = listing()
    for mode in ((), ("--complete",)):
        r = gct("report", "all", "--format", ",".join(FORMATS), *mode, *at, ok=4)
        assert r["cause"] == f"Document {doc} was deleted but its removal did not finish.", r
    assert listing() == before
    gct("doc", "delete", doc, "--confirm", token, *at)
    folder = Path(gct("report", "all", "--format", ",".join(FORMATS), *at)["folder"])
    assert leakage.found(folder, lw["seeds"]) == {}


def test_missing_evidence_in_both_modes_of_every_report_output(lw):
    """Missing evidence (K8, K9): a private document's documents-checked entry is in the complete gap report only,
    and the sanitised one says private documents are not listed in the privacy notice's words, with no count; an
    expected item declaring a private document is withheld in sanitised mode (its seeds are the matrix's LEAK-12 and
    LEAK-13); the three missing states print verbatim in both modes of every format's gap report."""
    from gct.stages import missing

    c, s = exported(lw, "complete"), exported(lw, "sanitised")
    private = {x["id"] for x in c["documents"] if x["label"] == "private"}
    assert [x for x in c["gap_report"]["documents_checked"] if x["id"] in private]
    assert {x["id"] for x in s["gap_report"]["documents_checked"]} <= {x["id"] for x in s["documents"]}
    assert dataset.PRIVACY_NOTICE in s["notices"] and not [n for n in s["notices"] if re.search(r"\d", n)]
    assert len(s["gap_report"]["expectations"]) < len(c["gap_report"]["expectations"])
    for mode in dataset.MODES:
        d = exported(lw, mode)
        assert {missing.INCOMPLETE, missing.NOT_FOUND_IN, missing.CONFIRMED} <= {
            x["state"] for x in d["gap_report"]["expectations"]}, mode
        for fmt in FORMATS:
            stem = "gap-report"
            paths = [p for p in files_of(lw[mode], fmt) if fmt == "json" or stem in p.as_posix()]
            for state in (missing.INCOMPLETE, missing.NOT_FOUND_IN, missing.CONFIRMED):
                assert has(paths, state), (mode, fmt, state)


def test_the_k1_claim_shows_its_recomputed_review_state_in_every_sanitised_output(lw):
    """Review state of a shareable claim (K1, recompute): the claim 7.1's disagreement and raised flag sit on shows,
    in sanitised mode, a review text computed from shipped records only (no withheld marker), printed in every
    format's own K1 row (read back, selected by claim_id); that it equals sanitised(W') is the differential's
    (test_sanitised_w_and_w_prime_render_...)."""
    c, s = exported(lw, "complete"), exported(lw, "sanitised")
    k1 = next(r for r in c["evidence_index"] if r["label"] == "shareable" and "disagreement" in r["review_text"]
              and "flags raised" in r["review_text"])
    row = next(r for r in s["evidence_index"] if r["claim_id"] == k1["claim_id"])
    assert row["review_text"] != k1["review_text"] and "review_text" not in row["withheld"]
    for fmt in FORMATS:  # the K1 row itself as each format reads it back (Codex review F3), not any text of the file
        m = module(fmt)
        claim = m.printed(row["claim_id"], row, "claim_id", fmt)
        printed = [r for r in back(fmt, lw["sanitised"])["rows"]["evidence_index"] if r["claim_id"] == claim]
        assert [r["review_text"] for r in printed] == [m.printed(x["review_text"], x, "review_text", fmt) for x in
                                                        s["evidence_index"] if x["claim_id"] == row["claim_id"]], fmt


def test_renewed_review_and_corrections_follow_their_labels(lw):
    """Renewed review (K6, C7, L4): a needs-re-review entry whose target is withheld is not in sanitised mode; L4's
    replaced values of a private claim are complete only (the matrix's LEAK-01 to LEAK-03); a shareable claim's
    correction keeps its values in sanitised mode with its contributor, organisation, capacity and rationale
    withheld (dataset.WITHHELD), printed as withheld in every printed format."""
    c, s = exported(lw, "complete"), exported(lw, "sanitised")
    kinds = set(dataset.LEDGER_KINDS.values())

    def named(d):  # the review records the entries name: "kind #id on version V( with claim C version W)"
        return {(m[1], int(m[2])) for r in d["evidence_index"] for e in r["needs_renewed_review"]
                if (m := ENTRY.match(e))[1] in kinds}

    def shipped(d):
        return {(x["kind"], x["id"]) for r in d["evidence_index"] for x in r["records"]["rows"]}

    assert named(c) - shipped(s)  # the check can fail: the review citing co3-leak.pdf (leakage_fixture, K6)
    assert named(s) and named(s) <= shipped(s)
    corrected = [r for r in s["evidence_index"] if r["correction"]]
    who = [f"correction.{k}" for k in ("contributor", "organisation", "capacity", "rationale")]
    assert corrected and all(set(who) <= set(r["withheld"]) and not any(r["correction"][k[11:]] for k in who)
                             for r in corrected)
    for fmt in FORMATS[1:]:  # the JSON export is the dataset read above (null and the withheld list)
        assert has(files_of(lw["sanitised"], fmt), module(fmt).printed(None, corrected[0], who[0], fmt)), fmt


def test_every_document_private_leaves_empty_sanitised_outputs(lw, tmp_path):
    """Empty results (P10): with every document private, every sanitised output holds no row, no document or company
    name and no count of withheld rows; it shows the data version, the empty-result line and the privacy notice; the
    preview sentence and the relabel command are in no output."""
    root = Path(shutil.copytree(lw["root"], tmp_path / "ws"))
    at = ("--workspace", root)
    docs = gct("doc", "list", *at)["documents"]
    for x in docs:
        if x["state"] != "deleted" and x["privacy"] == "shareable":
            gct("doc", "label", x["id"], "--privacy", "private", "--as", "kim", "--capacity", f.CAPACITY,
                "--rationale", "not for the client", *at)
    gct("publish", "build", *at)
    folder = Path(gct("report", "all", "--format", ",".join(FORMATS), *at)["folder"])
    d = json.loads((folder / "dataset.json").read_text(encoding="utf-8"))
    names = {x["logical_name"] for x in docs if x["logical_name"]} | {x["company"] for x in docs}
    assert dataset.sanitised_empty(d["data_version"]["version"]) in d["notices"]
    for fmt in FORMATS:
        paths = files_of(folder, fmt)
        assert not [t for t, rows in back(fmt, folder)["rows"].items() if rows], fmt
        assert not [n for n in names if has(paths, n)], fmt
        assert all(has(paths, n) for n in d["notices"]) and has(paths, dataset.PRIVACY_NOTICE), fmt
        assert not has(paths, dataset.ALL_PRIVATE) and not has(paths, dataset.RELABEL), fmt
    assert leakage.found(folder, lw["seeds"]) == {}


# AC 13: coverage. Content types, each mapped to seeds the matrix finds; every string leaf path of the complete
# dataset pinned to a seed or to "machine" with its reason (words.* and columns.* are GCT's own words and keys).
CONTENT = {
    "documents": {"DOC-NAME-01", "DOC-NAME-02", "DOC-COMPANY-01", "FILE-PDF-01", "LEAK-09", "LEAK-14"},
    "records": {"CON-NAME-01", "CON-CONCLUSION-01", "CON-NOTE-01", "CON-RATIONALE-01", "LEAK-21", "LEAK-22",
                "LEAK-29"},
    "graphs": {"DOC-QUOTE-01", "LEAK-08", "LEAK-23", "LEAK-24", "LEAK-25", "LEAK-26", "LEAK-27", "LEAK-28"},
    "indexes": {"DOC-METRIC-01", "DOC-CELL-01", "DOC-REASON-01", "LEAK-04", "LEAK-05", "LEAK-06", "LEAK-07"},
    "evidence images": {"IMG-PAGE-01", "IMG-MARKER-01"},
}
MACHINE = {"words", "columns"}  # prefixes
GCT_WORDS = "machine: GCT's own words"
CODE = "machine: a stored code or GCT's word for it"
DATE = "machine: a date or time GCT writes"
ADDRESS = "machine: a cell address or sheet position"
IDS = "machine: an id or a hash of ids"
SAME = "related_claims: one builder with evidence_index (test_related_claims_are_evidence_index_rows_of_one_builder)"
UNIT = ("machine: a unit key, or printed:<the unit text folded to lower-case letters and digits>, which no seed"
        " survives; the printed unit text itself is evidence_index[].unit_text (LEAK-05)")
COVERAGE = {
    "built_at": DATE,
    "data_version.built_at": DATE,
    "data_version.history_mark": IDS,
    "data_version.sha256": IDS,
    "documents[].availability": CODE,
    "documents[].company": "DOC-COMPANY-01",
    "documents[].file_type": CODE,
    "documents[].kind": CODE,
    "documents[].label": CODE,
    "documents[].name": "LEAK-15",
    "documents[].state": CODE,
    "evidence[].cell": ADDRESS,
    "evidence[].header[]": "LEAK-11",
    "evidence[].kind": CODE,
    "evidence[].row[]": "DOC-REASON-01",
    "evidence[].sha256": "FILE-PDF-04",
    "evidence[].sheet": "LEAK-10",
    "evidence_index[].boundary": CODE,
    "evidence_index[].boundary_detail": "LEAK-07",
    "evidence_index[].boundary_state": CODE,
    "evidence_index[].cell": ADDRESS,
    "evidence_index[].claim_state": CODE,
    "evidence_index[].company": "DOC-COMPANY-01",
    "evidence_index[].correction.capacity": "CON-CAPACITY-01",
    "evidence_index[].correction.contributor": "CON-NAME-01",
    "evidence_index[].correction.created_at": DATE,
    "evidence_index[].correction.fields[]": CODE,
    "evidence_index[].correction.organisation": "CON-ORG-02",
    "evidence_index[].correction.rationale": "CON-CORRECTION-01",
    "evidence_index[].correction.replaced.metric_text": "DOC-CELL-01",
    "evidence_index[].correction.replaced.period_label": "LEAK-03",
    "evidence_index[].correction.replaced.unit_text": "LEAK-02",
    "evidence_index[].correction.replaced.value_text": "LEAK-01",
    "evidence_index[].disagreements.rows[].label": CODE,
    "evidence_index[].disagreements.rows[].text": GCT_WORDS,
    "evidence_index[].document_name": "LEAK-15",
    "evidence_index[].evidence_text": "DOC-REASON-01",
    "evidence_index[].label": CODE,
    "evidence_index[].match_reason": GCT_WORDS,
    "evidence_index[].match_state": CODE,
    "evidence_index[].metric_text": "DOC-METRIC-01",
    "evidence_index[].needs_renewed_review[]": GCT_WORDS,
    "evidence_index[].period_end": DATE,
    "evidence_index[].period_label": "LEAK-06",
    "evidence_index[].period_start": DATE,
    "evidence_index[].period_state": CODE,
    "evidence_index[].records.rows[].capacity": "CON-CAPACITY-01",
    "evidence_index[].records.rows[].conclusion": "CON-CONCLUSION-01",
    "evidence_index[].records.rows[].contributor": "CON-NAME-01",
    "evidence_index[].records.rows[].created_at": DATE,
    "evidence_index[].records.rows[].decision": CODE,
    "evidence_index[].records.rows[].earlier": CODE,
    "evidence_index[].records.rows[].evidence_note": "LEAK-22",
    "evidence_index[].records.rows[].evidence_state": CODE,
    "evidence_index[].records.rows[].kind": CODE,
    "evidence_index[].records.rows[].label": CODE,
    "evidence_index[].records.rows[].note": "CON-NOTE-01",
    "evidence_index[].records.rows[].organisation": "CON-ORG-02",
    "evidence_index[].records.rows[].outcome": CODE,
    "evidence_index[].records.rows[].rationale": "CON-RATIONALE-01",
    "evidence_index[].records.rows[].reviewed_on": DATE,
    "evidence_index[].records.rows[].scope": "LEAK-21",
    "evidence_index[].review_text": GCT_WORDS,
    "evidence_index[].sheet": "LEAK-10",
    "evidence_index[].unit_text": "LEAK-05",
    "evidence_index[].value_text": "LEAK-04",
    "filter.text": "DOC-METRIC-01",
    "gap_report.archived[].archived_at": DATE,
    "gap_report.archived[].label": CODE,
    "gap_report.archived[].name": "DOC-NAME-02",
    "gap_report.candidates[].adjudications.rows[].capacity": "CON-CAPACITY-01",
    "gap_report.candidates[].adjudications.rows[].closure_reason": "LEAK-27",
    "gap_report.candidates[].adjudications.rows[].contributor": "CON-NAME-01",
    "gap_report.candidates[].adjudications.rows[].counter_note": "LEAK-25",
    "gap_report.candidates[].adjudications.rows[].created_at": DATE,
    "gap_report.candidates[].adjudications.rows[].earlier": CODE,
    "gap_report.candidates[].adjudications.rows[].follow_up": "LEAK-26",
    "gap_report.candidates[].adjudications.rows[].label": CODE,
    "gap_report.candidates[].adjudications.rows[].organisation": "CON-ORG-02",
    "gap_report.candidates[].adjudications.rows[].procedure": "LEAK-24",
    "gap_report.candidates[].adjudications.rows[].rationale": "LEAK-28",
    "gap_report.candidates[].adjudications.rows[].state_text": GCT_WORDS,
    "gap_report.candidates[].label": CODE,
    "gap_report.candidates[].metric_key": CODE,
    "gap_report.candidates[].period_end": DATE,
    "gap_report.candidates[].period_start": DATE,
    "gap_report.candidates[].reporting_boundary": CODE,
    "gap_report.candidates[].restatements.rows[].capacity": "CON-CAPACITY-01",
    "gap_report.candidates[].restatements.rows[].contributor": "CON-NAME-01",
    "gap_report.candidates[].restatements.rows[].created_at": DATE,
    "gap_report.candidates[].restatements.rows[].decision": CODE,
    "gap_report.candidates[].restatements.rows[].earlier": CODE,
    "gap_report.candidates[].restatements.rows[].evidence_quote": "LEAK-08",
    "gap_report.candidates[].restatements.rows[].label": CODE,
    "gap_report.candidates[].restatements.rows[].organisation": "CON-ORG-02",
    "gap_report.candidates[].restatements.rows[].rationale": "LEAK-23",
    "gap_report.candidates[].series_id": IDS,
    "gap_report.candidates[].shown_state": CODE,
    "gap_report.candidates[].state": CODE,
    "gap_report.candidates[].unit_family": UNIT,
    "gap_report.deleted[].deleted_at": DATE,
    "gap_report.deleted[].label": CODE,
    "gap_report.documents_checked[].label": CODE,
    "gap_report.documents_checked[].limit_note": GCT_WORDS,
    "gap_report.expectations[].accepted.rows[].label": CODE,
    "gap_report.expectations[].accepted.rows[].type": CODE,
    "gap_report.expectations[].company": "DOC-COMPANY-01",
    "gap_report.expectations[].confirmation.capacity": "CON-CAPACITY-01",
    "gap_report.expectations[].confirmation.contributor": "CON-NAME-01",
    "gap_report.expectations[].confirmation.created_at": DATE,
    "gap_report.expectations[].confirmation.organisation": "CON-ORG-02",
    "gap_report.expectations[].confirmation.rationale": "LEAK-29",
    "gap_report.expectations[].declared.rows[].label": CODE,
    "gap_report.expectations[].declared.rows[].state": CODE,
    "gap_report.expectations[].label": CODE,
    "gap_report.expectations[].metric_text": "LEAK-12",
    "gap_report.expectations[].period_label": "LEAK-13",
    "gap_report.expectations[].reasons.rows[].label": CODE,
    "gap_report.expectations[].reasons.rows[].reason": GCT_WORDS,
    "gap_report.expectations[].reasons.rows[].type": CODE,
    "gap_report.expectations[].state": CODE,
    "gap_report.extraction_limits[].label": CODE,
    "gap_report.extraction_limits[].reason": GCT_WORDS,
    "gap_report.extraction_limits[].type": CODE,
    "gct_version": GCT_WORDS,
    "mode": CODE,
    "notices[]": GCT_WORDS,
    "related_claims[].boundary": SAME,
    "related_claims[].boundary_detail": SAME,
    "related_claims[].boundary_state": SAME,
    "related_claims[].claim_state": SAME,
    "related_claims[].company": "DOC-COMPANY-01",
    "related_claims[].correction.capacity": SAME,
    "related_claims[].correction.contributor": SAME,
    "related_claims[].correction.created_at": SAME,
    "related_claims[].correction.fields[]": SAME,
    "related_claims[].correction.organisation": SAME,
    "related_claims[].correction.rationale": SAME,
    "related_claims[].correction.replaced.metric_text": SAME,
    "related_claims[].correction.replaced.period_label": SAME,
    "related_claims[].correction.replaced.unit_text": SAME,
    "related_claims[].correction.replaced.value_text": SAME,
    "related_claims[].document_name": "LEAK-09",
    "related_claims[].evidence_text": SAME,
    "related_claims[].label": SAME,
    "related_claims[].match_reason": SAME,
    "related_claims[].match_state": SAME,
    "related_claims[].metric_text": SAME,
    "related_claims[].period_end": SAME,
    "related_claims[].period_label": SAME,
    "related_claims[].period_start": SAME,
    "related_claims[].period_state": SAME,
    "related_claims[].review_text": SAME,
    "related_claims[].unit_text": SAME,
    "related_claims[].value_text": SAME,
    "restatement_log[].adjudications.rows[].capacity": "CON-CAPACITY-01",
    "restatement_log[].adjudications.rows[].closure_reason": "LEAK-27",
    "restatement_log[].adjudications.rows[].contributor": "CON-NAME-01",
    "restatement_log[].adjudications.rows[].counter_note": "LEAK-25",
    "restatement_log[].adjudications.rows[].created_at": DATE,
    "restatement_log[].adjudications.rows[].earlier": CODE,
    "restatement_log[].adjudications.rows[].follow_up": "LEAK-26",
    "restatement_log[].adjudications.rows[].label": CODE,
    "restatement_log[].adjudications.rows[].organisation": "CON-ORG-02",
    "restatement_log[].adjudications.rows[].procedure": "LEAK-24",
    "restatement_log[].adjudications.rows[].rationale": "LEAK-28",
    "restatement_log[].adjudications.rows[].state_text": GCT_WORDS,
    "restatement_log[].label": CODE,
    "restatement_log[].metric_key": CODE,
    "restatement_log[].period_end": DATE,
    "restatement_log[].period_start": DATE,
    "restatement_log[].reason": GCT_WORDS,
    "restatement_log[].reporting_boundary": CODE,
    "restatement_log[].restatements.rows[].capacity": "CON-CAPACITY-01",
    "restatement_log[].restatements.rows[].contributor": "CON-NAME-01",
    "restatement_log[].restatements.rows[].created_at": DATE,
    "restatement_log[].restatements.rows[].decision": CODE,
    "restatement_log[].restatements.rows[].earlier": CODE,
    "restatement_log[].restatements.rows[].evidence_quote": "LEAK-08",
    "restatement_log[].restatements.rows[].label": CODE,
    "restatement_log[].restatements.rows[].organisation": "CON-ORG-02",
    "restatement_log[].restatements.rows[].rationale": "LEAK-23",
    "restatement_log[].series_id": IDS,
    "restatement_log[].shown_state": CODE,
    "restatement_log[].state": CODE,
    "restatement_log[].unit_family": UNIT,
    "unpaired_restatements.rows[].capacity": "CON-CAPACITY-01",
    "unpaired_restatements.rows[].contributor": "CON-NAME-01",
    "unpaired_restatements.rows[].created_at": DATE,
    "unpaired_restatements.rows[].decision": CODE,
    "unpaired_restatements.rows[].earlier": CODE,
    "unpaired_restatements.rows[].evidence_quote": "LEAK-08",
    "unpaired_restatements.rows[].label": CODE,
    "unpaired_restatements.rows[].organisation": "CON-ORG-02",
    "unpaired_restatements.rows[].rationale": "LEAK-23",
}


def leaves(x, path=""):
    """(dotted path, string) of every string leaf; a sub-table's rows are "<key>.rows[]"."""
    if isinstance(x, dict):
        if set(x) == {"id", "title", "rows", "parent_key", "columns"}:
            for r in x["rows"]:
                yield from leaves(r, f"{path}.rows[]")
            return
        for k, v in x.items():
            yield from leaves(v, f"{path}.{k}" if path else k)
    elif isinstance(x, list):
        for v in x:
            yield from leaves(v, f"{path}[]")
    elif isinstance(x, str):
        yield path, x


def coverage_builds(root) -> list:
    """The complete dataset of the whole version, of a selection whose pairs have sides outside it (co3-2025.pdf,
    acme-2025.pdf and the private Acme report: co3-leak.pdf's claims among the related ones) and of a text filter
    (the filter echo)."""
    d, _ = build(root, mode="complete")
    docs = [x["id"] for x in d["documents"] if x["name"] in ("co3-2025.pdf", "acme-2025.pdf", f.NAME)]
    return [d, build(root, mode="complete", filter={"document": docs})[0],
            build(root, mode="complete", filter={"text": f.METRIC})[0]]


def test_every_content_type_maps_to_seeds_the_outputs_find(lw):
    """AC 13: documents, records, graphs (restatement-log pairs and comparisons; M1 has no series chart), indexes
    (evidence index rows, documents checked; SQLite index pages through the packages' raw scan) and evidence images,
    each to seeds found in a complete output of some format (the matrix); calculated values have no row in M1."""
    for kind, names in CONTENT.items():
        for n in names:
            assert any(n in lw["hits"][fmt, "complete"] for fmt in FORMATS), (kind, n)
        assert any(n in lw["hits"]["complete package"] for n in names) or kind == "evidence images", kind


def test_every_string_leaf_of_the_complete_dataset_is_pinned_to_a_seed_or_machine(lw):
    """AC 13 (7.7-P3): every string leaf path of the complete dataset is in COVERAGE, so a new field fails until it is
    mapped; a path mapped to a seed holds that seed's value in some leaf; a machine leaf (words.* and columns.*
    included) holds no seed. related_claims[] paths are covered through one builder (the next test)."""
    values = {s["name"]: s["value"] for s in lf.SEEDS if s["value"] and s["kind"] != "image"}
    at = {}
    for d in coverage_builds(lw["root"]):
        for p, v in leaves(d):
            at.setdefault(p, []).append(v)
    seen = {p for p in at if p.split(".")[0] not in MACHINE}
    assert seen == set(COVERAGE), (sorted(seen - set(COVERAGE)), sorted(set(COVERAGE) - seen))
    for p, vs in at.items():
        want = COVERAGE.get(p, GCT_WORDS)
        if want in values:
            assert any(values[want] in v for v in vs), (p, want)
        elif want != SAME:
            assert not [n for n, x in values.items() for v in vs if x in v], p


def cites(d, ids) -> list:
    """The evidence entries ids name, without their ids (evidence ids number each dataset's own entries)."""
    return [{k: v for k, v in next(e for e in d["evidence"] if e["id"] == i).items() if k != "id"} for i in ids]


def test_related_claims_are_evidence_index_rows_of_one_builder(lw):
    """AC 13, Moka's ruling: related_claims[].X is covered by evidence_index[].X through one builder. In both modes,
    each related claim of acme-2025.pdf's selection (and of co3-2025.pdf's in complete mode) equals that claim's
    evidence-index row in its own document's selection on every key, withheld markers included, its evidence
    compared by the entries it cites; both directions of the key sets come from the dataset, so a field added to
    one row and not the other fails."""
    root, checked = lw["root"], 0
    for mode in dataset.MODES:
        d, _ = build(root, mode=mode)
        names = {x["name"]: x["id"] for x in d["documents"]}
        for doc in ["acme-2025.pdf"] + (["co3-2025.pdf"] if mode == "complete" else []):
            a, _ = build(root, mode=mode, filter={"document": [names[doc]]})
            assert a["related_claims"], (mode, doc)
            for x in a["related_claims"]:
                b, _ = build(root, mode=mode, filter={"document": [x["document_id"]]})
                row = next(r for r in b["evidence_index"] if r["claim_id"] == x["claim_id"])
                assert set(x) == set(row), (mode, sorted(set(x) ^ set(row)))
                assert cites(a, x["evidence"]) == cites(b, row["evidence"]), (mode, x["claim_id"])
                assert {k: v for k, v in x.items() if k != "evidence"} == {
                    k: v for k, v in row.items() if k != "evidence"}, (mode, x["claim_id"])
                checked += 1
    assert checked >= 3


# The scanner on xref entries proven free (Moka's AC 17 finding: a real PDF's xref stream lists free ids inside its
# /Index ranges), on synthetic PDFs only.

PAGE = {1: b"<< /Type /Catalog /Pages 2 0 R >>", 2: b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        3: b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 100 100] >>"}


def xref_pdf(*revisions) -> bytes:
    """A PDF of one or more revisions, each (objects {id: body}, /Index ranges, entry types {id: 1, "missing" for an
    in-use entry pointing past the end of the file, or (2, stream id, index) for an object stream member}), indexed
    by an xref stream (W [1 4 2]) with /Prev to the
    one before; an id of a range not in the types is free (type 0). Offsets are those of the newest copy so far."""
    out, at, prev = bytearray(b"%PDF-1.5\n"), {}, None
    for n, (objects, ranges, kinds) in enumerate(revisions):
        for i, body in objects.items():
            at[i] = len(out)
            out += b"%d 0 obj\n%s\nendobj\n" % (i, body)
        rows = b""
        for start, count in ranges:
            for i in range(start, start + count):
                kind = kinds.get(i, 0)
                if isinstance(kind, tuple):  # (2, object stream id, index): a member of an object stream
                    rows += bytes([2]) + kind[1].to_bytes(4, "big") + kind[2].to_bytes(2, "big")
                    continue
                rows += bytes([kind != 0]) + (10 ** 6 if kind == "missing" else at.get(i, 0) * (kind == 1)).to_bytes(
                    4, "big") + b"\0\0"
        pos = len(out)
        out += b"%d 0 obj\n<< /Type /XRef /Size %d /Index [%s] /W [1 4 2] /Root 1 0 R /Length %d%s >>\nstream\n" % (
            90 + n, max(s + c for s, c in ranges), b" ".join(b"%d %d" % r for r in ranges), len(rows),
            b"" if prev is None else b" /Prev %d" % prev)
        out += rows + b"\nendstream\nendobj\nstartxref\n%d\n%%%%EOF\n" % pos
        prev = pos
    return bytes(out)


def seed(n) -> dict:
    return {"name": f"FREE-{n}", "kind": "document_text", "expect": "never", "value": f"GCT-LEAK-SCAN-FREE-{n:02d}"}


IN_USE = {1: 1, 2: 1, 3: 1, 4: 1}


def test_free_entries_inside_an_xref_streams_index_ranges_are_skipped(tmp_path):
    """An xref stream lists ids 10 and 11 free inside its second /Index range (pdfminer's get_objids lists 11, reading
    the first range's entry 1): no header and no object stream holds them, so found scans with {} and no Unreadable,
    and a seed in object 12 beside them, in use, is found in that object."""
    p = tmp_path / "free.pdf"
    p.write_bytes(xref_pdf(({**PAGE, 4: b"(x)"}, [(0, 5), (10, 2)], IN_USE)))
    assert leakage.found(p, [seed(1)]) == {}
    p.write_bytes(xref_pdf(({**PAGE, 4: b"(x)", 12: b"(%s)" % seed(2)["value"].encode()}, [(0, 5), (10, 3)],
                            {**IN_USE, 12: 1})))
    assert "free.pdf: object 12" in leakage.found(p, [seed(2)])["FREE-2"]


def test_in_use_ids_pdfminer_leaves_unlisted_are_still_read(tmp_path):
    """Moka's check on the same pdfminer reading: /Index [0 6 10 6] with ids 0, 4 and 5 free, so get_objids reads
    ids 10, 14 and 15 against the stream's entries 0, 4 and 5 (free) and leaves them out, and lists 11 to 13, free,
    against ids 1 to 3's entries. Object 14 (in use, a seed) and object 10 (an object stream, in use) whose member 15
    holds a seed are still read, by their headers and the member parse; 11 to 13 are proven free. No Unreadable."""
    member = b"15 0 (%s)" % seed(7)["value"].encode()
    objects = {**PAGE, 10: b"<< /Type /ObjStm /N 1 /First 5 /Length %d >>\nstream\n%s\nendstream" % (
        len(member), member), 14: b"(%s)" % seed(6)["value"].encode()}
    p = tmp_path / "unlisted.pdf"
    p.write_bytes(xref_pdf((objects, [(0, 6), (10, 6)], {1: 1, 2: 1, 3: 1, 10: 1, 14: 1, 15: (2, 10, 0)})))
    import pdfplumber

    with pdfplumber.open(p) as pdf:  # the reading the test is about: pdfminer lists none of the three
        assert not {10, 14, 15} & {i for x in pdf.doc.xrefs for i in x.get_objids()}
    hit = leakage.found(p, [seed(6), seed(7)])
    assert [x for x in hit["FREE-6"] if "object 14" in x] and [x for x in hit["FREE-7"] if "member 15" in x], hit


def test_a_free_entry_with_a_header_or_in_an_object_stream_is_unreadable(tmp_path):
    """Object 11 marked free in the second /Index range, but with an "11 0 obj" header in the file, or a member of an
    object stream: not proven free, so found raises Unreadable naming object 11 (its seed is still found)."""
    p = tmp_path / "held.pdf"
    member = b"11 0 (%s)" % seed(5)["value"].encode()
    for extra in ({11: b"(%s)" % seed(5)["value"].encode()},
                  {5: b"<< /Type /ObjStm /N 1 /First 5 /Length %d >>\nstream\n%s\nendstream" % (len(member), member)}):
        p.write_bytes(xref_pdf(({**PAGE, 4: b"(x)", **extra}, [(0, 6), (10, 2)], {**IN_USE, 5: int(5 in extra)})))
        with pytest.raises(leakage.Unreadable) as e:
            leakage.found(p, [seed(5)])
        assert [x for x in e.value.places if x.startswith("held.pdf: object 11 ")] == [
            x for x in e.value.places], e.value.places
        assert "FREE-5" in e.value.found


def test_an_in_use_entry_pointing_at_nothing_is_unreadable(tmp_path):
    """An xref stream entry in use (type 1) whose offset points past the end of the file: found raises Unreadable
    naming that object."""
    p = tmp_path / "missing.pdf"
    p.write_bytes(xref_pdf((PAGE, [(0, 5)], {**IN_USE, 4: "missing"})))
    with pytest.raises(leakage.Unreadable) as e:
        leakage.found(p, [seed(3)])
    assert [x for x in e.value.places if x.startswith("missing.pdf: object 4 ")], e.value.places


def test_an_id_free_in_one_revision_and_in_use_in_another_is_read(tmp_path):
    """Object 11 in use in the first revision and free in the update's second /Index range: not proven free, so it is
    read: its seed is found in it, and with its in-use entry pointing at nothing found raises Unreadable naming it."""
    p = tmp_path / "mixed.pdf"
    update = ({4: b"(x)"}, [(0, 5), (10, 2)], IN_USE)
    p.write_bytes(xref_pdf(({**PAGE, 11: b"(%s)" % seed(4)["value"].encode()}, [(0, 4), (11, 1)],
                            {1: 1, 2: 1, 3: 1, 11: 1}), update))
    assert "mixed.pdf: object 11" in leakage.found(p, [seed(4)])["FREE-4"]
    p.write_bytes(xref_pdf((PAGE, [(0, 4), (11, 1)], {1: 1, 2: 1, 3: 1, 11: "missing"}), update))
    with pytest.raises(leakage.Unreadable) as e:
        leakage.found(p, [seed(4)])
    assert [x for x in e.value.places if x.startswith("mixed.pdf: object 11 ")], e.value.places


# The scanner on Indexed images (Moka's AC 17 finding: real reports' Indexed images were unread), synthetic only.

def indexed_pdf(base: bytes, palette: list, bits: int, lookup: str, size=(10, 10)) -> tuple[bytes, list]:
    """A PDF holding one image XObject in [/Indexed base hival lookup] (lookup "string" a hex string in the array,
    "stream" an indirect stream; base "/DeviceRGB" or an indirect reference to "/DeviceCMYK" when it is b"ref"),
    pixel (x, y) the palette entry (x + y) % len(palette), rows padded to a byte; and the indices it draws."""
    w, h = size
    indices = [(x + y) % len(palette) for y in range(h) for x in range(w)]
    rows = b""
    for y in range(h):
        row = "".join(format(i, f"0{bits}b") for i in indices[y * w:(y + 1) * w])
        rows += int(row.ljust(-(-len(row) // 8) * 8, "0"), 2).to_bytes(-(-len(row) // 8), "big")
    table = bytes(v for entry in palette for v in entry)
    objects = {**PAGE}
    if base == b"ref":
        objects[6], base = b"/DeviceCMYK", b"6 0 R"
    if lookup == "stream":
        objects[5], ref = b"<< /Length %d >>\nstream\n%s\nendstream" % (len(table), table), b"5 0 R"
    else:
        ref = b"<%s>" % table.hex().encode()
    objects[4] = b"<< /Type /XObject /Subtype /Image /Width %d /Height %d /BitsPerComponent %d /ColorSpace " \
                 b"[/Indexed %s %d %s] /Length %d >>\nstream\n%s\nendstream" % (
                     w, h, bits, base, len(palette) - 1, ref, len(rows), rows)
    return xref_pdf((objects, [(0, max(objects) + 1)], dict.fromkeys(objects, 1))), indices


@pytest.mark.parametrize("base, bits, lookup", [(b"/DeviceRGB", 8, "string"), (b"ref", 4, "stream"),
                                                (b"/DeviceRGB", 2, "stream"), (b"ref", 1, "string")])
def test_an_indexed_image_is_found_by_the_pixel_hash_of_its_palette_colours(tmp_path, base, bits, lookup):
    """An Indexed image over DeviceRGB or a referenced DeviceCMYK, its lookup a string or a stream, at 8, 4, 2 and 1
    bits: found reads it as the direct image of its palette colours, so a seed of that image's pixel hash is found."""
    n = 3 if base == b"/DeviceRGB" else 4
    palette = [bytes((37 * i + 11 * k) % 256 for k in range(n)) for i in range(min(2 ** bits, 6))]
    data, indices = indexed_pdf(base, palette, bits, lookup)
    image = Image.frombytes("RGB" if n == 3 else "CMYK", (10, 10), b"".join(palette[i] for i in indices))
    p = tmp_path / "indexed.pdf"
    p.write_bytes(data)
    value = hashlib.sha256(image.convert("RGB").tobytes()).hexdigest()
    hit = leakage.found(p, [{"name": "IDX", "kind": "image", "expect": "never", "value": value}])
    assert hit == {"IDX": ["indexed.pdf: object 4 image: pixel hash"]}


def test_a_marker_coloured_indexed_image_is_found_under_the_marker_seed(tmp_path):
    """An 8-bit Indexed RGB image whose palette holds the marker colour: found under the value-None image seed."""
    p = tmp_path / "marker.pdf"
    p.write_bytes(indexed_pdf(b"/DeviceRGB", [bytes(leakage.MARKER)], 8, "string")[0])
    hit = leakage.found(p, [{"name": "MARK", "kind": "image", "expect": "never", "value": None}])
    assert hit == {"MARK": ["marker.pdf: object 4 image: marker colour"]}


def decode_pdf(colour_space: bytes, bits: int, samples: bytes, decode: bytes | None) -> bytes:
    """A 10 by 10 image XObject (object 4) drawn on the page, with an optional /Decode array."""
    objects = {**PAGE, 3: b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 100 100] /Resources << /XObject << /Im 4 0 R"
                          b" >> >> /Contents 5 0 R >>",
               4: b"<< /Type /XObject /Subtype /Image /Width 10 /Height 10 /BitsPerComponent %d /ColorSpace %s%s"
                  b" /Length %d >>\nstream\n%s\nendstream" % (
                      bits, colour_space, b"" if decode is None else b" /Decode " + decode, len(samples), samples),
               5: b"<< /Length 29 >>\nstream\nq 100 0 0 100 0 0 cm /Im Do Q\nendstream"}
    return xref_pdf((objects, [(0, 6)], dict.fromkeys(objects, 1)))


MARKED = [{"name": "MARK", "kind": "image", "expect": "never", "value": None},
          {"name": "SOLID", "kind": "image", "expect": "never",
           "value": hashlib.sha256(Image.new("RGB", (10, 10), leakage.MARKER).tobytes()).hexdigest()}]
BLACK = {"name": "BLACK", "kind": "image", "expect": "never",
         "value": hashlib.sha256(Image.new("RGB", (10, 10), (0, 0, 0)).tobytes()).hexdigest()}
PALETTE = b"[/Indexed /DeviceRGB 1 <000000%s>]" % bytes(leakage.MARKER).hex().encode()
PALETTE4 = b"[/Indexed /DeviceRGB 3 <000000000000000000%s>]" % bytes(leakage.MARKER).hex().encode()  # entry 3 marks


@pytest.mark.parametrize("colour_space, bits, samples, decode, seen", [
    (PALETTE, 1, b"\0" * 20, b"[1 0]", "marker"),  # Codex F2: zero samples decode to index 1, the marker
    (PALETTE, 1, b"\0" * 20, b"[0 1]", "black"),  # the default Decode, given: zero samples are index 0
    (PALETTE, 1, b"\0" * 20, None, "black"),  # no Decode
    (PALETTE4, 2, b"\0" * 30, b"[3 0]", "marker"),  # reversed, 2 bits: sample s to index 3 - s, 0 to 3
    (PALETTE, 8, b"\0" * 100, b"[1 1]", "marker")],  # every sample to index 1
    ids=["1bit-inverted", "1bit-default-given", "1bit-no-decode", "2bit-reversed", "8bit-constant"])
def test_an_indexed_images_decode_array_maps_samples_to_indices_before_the_palette(tmp_path, colour_space, bits,
                                                                                   samples, decode, seen):
    """Codex review F2: an Indexed image's /Decode maps each sample to its index (Dmin + s (Dmax - Dmin) / (2^bits -
    1)) before the palette lookup, so a [1 0] Decode over zero samples draws palette entry 1, the marker, and is
    found under the marker seed and by its pixel hash; the default Decode, given or not, draws entry 0 (black)."""
    p = tmp_path / "decode.pdf"
    p.write_bytes(decode_pdf(colour_space, bits, samples, decode))
    hit = leakage.found(p, [*MARKED, BLACK])
    where = "decode.pdf: object 4 image"
    assert hit == ({"MARK": [f"{where}: marker colour"], "SOLID": [f"{where}: pixel hash"]} if seen == "marker"
                   else {"BLACK": [f"{where}: pixel hash"]}), hit


def test_a_device_images_decode_array_is_applied(tmp_path):
    """The same rule for a direct image: 8-bit DeviceRGB samples 255, 55, 135 with /Decode [1 0 1 0 1 0] draw the
    marker colour (0, 200, 120), found under the marker seed and by its pixel hash."""
    p = tmp_path / "decode.pdf"
    p.write_bytes(decode_pdf(b"/DeviceRGB", 8, bytes((255, 55, 135)) * 100, b"[1 0 1 0 1 0]"))
    where = "decode.pdf: object 4 image"
    assert leakage.found(p, MARKED) == {"MARK": [f"{where}: marker colour"], "SOLID": [f"{where}: pixel hash"]}


def test_a_jpeg_images_decode_array_is_read_only_when_it_is_the_default(tmp_path):
    """DCT data: the default Decode [0 1 0 1 0 1] is read as Pillow decodes it; [1 0 1 0 1 0] is Unreadable (not
    applied on top of Pillow's reading)."""
    jpeg = io.BytesIO()
    Image.new("RGB", (10, 10), (40, 40, 40)).save(jpeg, "JPEG")
    jpeg = jpeg.getvalue()
    p = tmp_path / "jpeg.pdf"
    for decode, unreadable in ((b"[0 1 0 1 0 1]", False), (b"[1 0 1 0 1 0]", True)):
        p.write_bytes(decode_pdf(b"/DeviceRGB /Filter /DCTDecode", 8, jpeg, decode))
        if unreadable:
            with pytest.raises(leakage.Unreadable) as e:
                leakage.found(p, MARKED)
            assert [x.split(" (")[0] for x in e.value.places] == ["jpeg.pdf: object 4"], e.value.places
        else:
            assert leakage.found(p, MARKED) == {}


@pytest.mark.parametrize("colour_space, decode, bits", [
    (PALETTE, b"[0 1 0]", 1), (PALETTE, b"[0 2]", 1), (PALETTE, b"[/a /b]", 1), (b"/DeviceRGB", b"[1 0]", 8),
    (PALETTE, b"[0.5 0.5]", 1), (PALETTE, b"[1 0]", 2)],
    ids=["indexed-three-numbers", "indexed-above-hival", "indexed-names", "rgb-one-pair", "indexed-half-index",
         "indexed-fractional-2bit"])
def test_an_unsupported_decode_array_is_unreadable(tmp_path, colour_space, decode, bits):
    """A Decode of the wrong length, of names, or one that maps some sample to an index that is not whole or not in
    0 to hival (Codex fix check 1: [0.5 0.5] over zero samples, which Poppler draws as the marker and PDFium as
    black; [1 0] at 2 bits, whose samples 1 and 2 fall between indices): Unreadable naming the image, never a scan
    that passes."""
    samples = b"\0" * (20 if bits == 1 else 30) if colour_space == PALETTE else b"\0" * 300
    p = tmp_path / "bad.pdf"
    p.write_bytes(decode_pdf(colour_space, bits, samples, decode))
    with pytest.raises(leakage.Unreadable) as e:
        leakage.found(p, MARKED)
    assert [x.split(" (")[0] for x in e.value.places] == ["bad.pdf: object 4"], e.value.places
    assert "Decode" in e.value.places[0]


def test_an_inline_image_with_an_empty_filter_array_is_read(tmp_path):
    """Moka's AC 17 IndexError: an inline image whose dictionary holds /F [] (no filter), in a page's content and in
    a form XObject. pdfminer's content parser indexes the empty array; found reads both images, by pixel hash."""
    pixels = Image.new("RGB", (4, 4), (10, 20, 30)).tobytes()
    inline = b"q 4 0 0 4 0 0 cm BI /W 4 /H 4 /BPC 8 /CS /RGB /F [] ID %s EI Q" % pixels
    objects = {**PAGE, 3: b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 100 100] /Contents 5 0 R >>",
               4: b"<< /Type /XObject /Subtype /Form /BBox [0 0 10 10] /Length %d >>\nstream\n%s\nendstream" % (
                   len(inline), inline),
               5: b"<< /Length %d >>\nstream\n%s\nendstream" % (len(inline), inline)}
    p = tmp_path / "inline.pdf"
    p.write_bytes(xref_pdf((objects, [(0, 6)], dict.fromkeys(objects, 1))))
    value = hashlib.sha256(pixels).hexdigest()
    hit = leakage.found(p, [{"name": "INLINE", "kind": "image", "expect": "never", "value": value}])
    assert hit == {"INLINE": ["inline.pdf: object 4 form inline image 1: pixel hash",
                              "inline.pdf: page 1 inline image 1: pixel hash"]}
