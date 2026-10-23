"""Story 7.3: the XLSX and CSV formats (gct/reports/render_sheets.py) on tests/report_fixture.py. read_back(path) and
printed(value, row, key, fmt) are this format pair's part of the shared read-back contract (readiness resolution 3 as
amended) that story 7.7 calls for every format. Files are opened the way an outsider would: openpyxl.load_workbook,
GCT's own stdlib reader sheet._xlsx, zipfile on every member, and csv.reader with the excel dialect."""
import ast
import copy
import csv
import hashlib
import json
import re
import shutil
import sqlite3
import sys
import time
import zipfile
from pathlib import Path
from xml.etree import ElementTree

import pytest
from openpyxl import load_workbook

import report_fixture as f
from leakage import found
from test_dataset import renderer_problems
from test_gate_privacy_exports import build, env, w  # noqa: F401 (module fixtures: the fixture and its environment)
from test_messages import INTERNAL_WORDS, internal_names
from test_review import WORDS as GUARD_WORDS
from test_view import dump
from gct.reports import dataset, export, render_sheets as rs
from gct.stages import adjudicate, missing, sheet
from gct.workspace import labels, layout, registry, store

ROOT = Path(__file__).resolve().parent.parent
WORDS = dataset.words()
TITLES = WORDS["fixed"]["titles"]
REPORTS = ("evidence_index", "restatement_log", "gap_report")
SECTIONS = ("candidates", "expectations", "documents_checked", "extraction_limits", "archived", "deleted")
DOCS = {"document": [8, 1, 19]}  # a document filter, so pairs have sides outside the selection (related claims)
# The tables each report's output holds, in sheet order (7.3-P3, launch rulings B6 and B7), written out here.
SHEETS = {
    "evidence_index": ["evidence_index", "records", "disagreements", "documents"],
    "restatement_log": ["restatement_log", "restatements", "adjudications", "unpaired_restatements", "pair_sides",
                        "documents"],
    "gap_report": ["candidates", "restatements", "adjudications", "pair_sides", "expectations", "reasons", "accepted",
                   "declared", "documents_checked", "extraction_limits", "archived", "deleted", "documents"],
}
SUB_OF = {"records": "evidence_index", "disagreements": "evidence_index", "reasons": "expectations",
          "accepted": "expectations", "declared": "expectations"}
GUARD = re.compile(GUARD_WORDS.pattern + "|contradiction", re.I)
WHO = ("--as", "ana", "--capacity", "analyst", "--rationale", "checked on the page")


# The read-back contract

def printed(value, row, key, fmt):
    """How XLSX and CSV print a dataset value (fmt "xlsx" or "csv"): render_sheets.printed with GCT's words."""
    return rs.printed(value, row, key, fmt, WORDS)


def _key_of(table, headings):
    by_heading = {c["heading"]: c["key"] for c in dataset.COLUMNS[table]}
    return [by_heading[h] for h in headings]


def _split(rows):
    """A table sheet's rows as (band lines, headings, data rows): the headings follow the first blank line."""
    blank = next(i for i, r in enumerate(rows) if not any(r))
    width = len([x for x in rows[blank + 1] if x])
    return rows[:blank], rows[blank + 1][:width], [r[:width] for r in rows[blank + 2:]]


def _tables_of(path):
    """{table id: (headings, data rows)} and the About rows of one <stem>.xlsx or <stem>/ folder, as printed text."""
    path, ids = Path(path), {v: k for k, v in TITLES.items()}
    out = {}
    if path.suffix == ".xlsx":
        for title, rows in sheet._xlsx(path):
            if ids[title] == "about":
                about = rows
            else:
                out[ids[title]] = _split(rows)[1:]
    else:
        for p in sorted(path.glob("*.csv")):
            rows = csv_rows(p)
            if p.stem == "about":
                about = rows
            else:
                out[p.stem] = (rows[0], rows[1:])
    return out, about


def read_back(path) -> dict:
    """{data_version: the identity the About sheet or about.csv holds, rows: {table: [{columns key: printed text}]}}
    of one <stem>.xlsx or one <stem>/ CSV folder: evidence_index, restatement_log (pair rows only) and
    gap_report.<section> for each section the output holds. Values are the cell texts (a number as the file holds
    it), an empty cell "" (sheet._xlsx and the csv module)."""
    out, about = _tables_of(path)
    identity = {v: k for k, v in WORDS["fixed"]["identity"].items()}
    rows = {}
    for t, (headings, data) in out.items():
        if t in ("evidence_index", "restatement_log") or t in SECTIONS:
            keys = _key_of(t, headings)
            rows[t if t in ("evidence_index", "restatement_log") else f"gap_report.{t}"] = [
                dict(zip(keys, r)) for r in data]
    order = ["evidence_index", "restatement_log", *(f"gap_report.{s}" for s in SECTIONS)]
    return {"data_version": {identity[r[0]]: r[1] for r in about[1:] if r[0] in identity},
            "rows": {k: rows[k] for k in order if k in rows}}


# Helpers

def csv_rows(path) -> list:
    """Every row of a CSV file. CSV has no cell limit, so the read lifts the csv module's field limit (131,072 by
    default) and restores the process limit after."""
    limit = csv.field_size_limit(sys.maxsize)
    try:
        with open(path, encoding="utf-8-sig", newline="") as fh:
            return list(csv.reader(fh, csv.excel))
    finally:
        csv.field_size_limit(limit)


def render(d, folder, reports=REPORTS, formats=("xlsx", "csv")) -> dict:
    """Each format's write into folder (made here, standing in for 7.2's .incoming- folder)."""
    folder.mkdir(parents=True, exist_ok=True)
    return {fmt: export.FORMATS[fmt]["write"](d, list(reports), folder, {}) for fmt in formats}


def stem(report):
    return report.replace("_", "-")


def expected(d, report) -> list[tuple]:
    """[(table id, columns, rows)] a report's output holds, from the dataset, in sheet order."""
    out = []
    for t in SHEETS[report]:
        cols = d["columns"]["evidence_index" if t == "pair_sides" else t]
        if t in ("evidence_index", "restatement_log"):
            rows = d[t]
        elif t in SECTIONS:
            rows = d["gap_report"][t]
        elif t == "documents":
            rows = d["documents"]
        elif t == "unpaired_restatements":
            rows = d["unpaired_restatements"]["rows"]
        elif t == "pair_sides":
            pairs = d["restatement_log"] if report == "restatement_log" else d["gap_report"]["candidates"]
            by_id = {r["id"]: r for r in d["evidence_index"] + d["related_claims"]}
            rows = [by_id[i] for i in dict.fromkeys(c for p in pairs for c in (p["claim_a"], p["claim_b"]))]
        else:
            parent = SUB_OF.get(t) or ("restatement_log" if report == "restatement_log" else "candidates")
            parents = d[parent] if parent in d else d["gap_report"][parent]
            rows = [x for p in parents for x in p[t]["rows"]]
        out.append((t, cols, rows))
    return out


def row_text(row, cols, fmt):
    return [printed(rs.dotted(row, c["key"]), row, c["key"], fmt) for c in cols]


def is_attributed(rows, cols):
    keys = {c["key"] for c in cols} & {"contributor", "correction.contributor", "confirmation.contributor"}
    return any(rs.dotted(r, k) is not None for r in rows for k in keys)


def worksheets(path) -> dict:
    with zipfile.ZipFile(path) as z:
        return {n: z.read(n) for n in z.namelist() if n.startswith("xl/worksheets/")}


def notes_counts(out) -> dict:
    """{fmt: [the numbers in its notes]}"""
    return {fmt: [int(re.search(r"\d+", n).group()) for n in o["notes"]] for fmt, o in out.items()}


@pytest.fixture(scope="module")
def outs(w, env, tmp_path_factory):
    """Both modes, the whole version and the document filter: {(mode, filter name): (dataset, folder, writes)}."""
    got = {}
    for mode in dataset.MODES:
        for name, flt in (("all", None), ("docs", DOCS)):
            d, _ = build(w, mode=mode, filter=flt)
            folder = tmp_path_factory.mktemp(f"{mode}-{name}")
            got[mode, name] = d, folder, render(d, folder)
    return got


@pytest.fixture
def ws(w, env, tmp_path):
    return Path(shutil.copytree(w, tmp_path / "ws"))


# AC 1: the module, the format table and the files

def test_two_formats_one_module_no_workspace_and_openpyxl_inside_the_writer():
    assert export.FORMATS["xlsx"] == {"write": rs.write_xlsx, "crops": False}
    assert export.FORMATS["csv"] == {"write": rs.write_csv, "crops": False}
    source = (ROOT / "gct" / "reports" / "render_sheets.py").read_text(encoding="utf-8")
    assert renderer_problems(source) == []
    top = [n for n in ast.parse(source).body if isinstance(n, (ast.Import, ast.ImportFrom))]
    assert sorted(a.name for n in top for a in n.names) == ["csv", "json", "re"]  # openpyxl only inside write_xlsx
    assert "preview" not in rs.write_xlsx.__code__.co_varnames and "preview" not in rs.write_csv.__code__.co_varnames


def test_all_writes_every_report_in_both_formats_into_one_folder(ws, gct):
    r = gct("report", "all", "--format", "xlsx,csv", "--workspace", ws, "--json")
    assert r.code == 0, r.out
    folder = Path(r.data["folder"])
    listed = sorted(p.relative_to(folder).as_posix() for p in folder.rglob("*") if p.is_file())
    assert listed == sorted(r.data["files"])
    want = {f"{stem(x)}.xlsx" for x in REPORTS} | {f"{stem(x)}/{t}.csv" for x in REPORTS for t in ["about", *SHEETS[x]]}
    assert set(listed) == want
    for x in REPORTS:
        assert registry.checked_file_type(folder / f"{stem(x)}.xlsx") == "xlsx"
    assert [o["format"] for o in r.data["outputs"]] == ["xlsx", "csv"]
    human = gct("report", "evidence-index", "--format", "csv", "--workspace", ws)
    assert human.code == 0 and "csv: 0 cells start with ' so that a spreadsheet program reads them as text." in human.out


# AC 2: tables

def test_the_nesting_rs_names_is_the_datasets_own(outs):
    """Launch ruling B7: the sub-tables render_sheets writes for a list are the ones its rows carry (seen on the
    fixture's rows, both modes), and every columns table but related_claims belongs to some report."""
    seen = {}
    for d, _, _ in outs.values():
        lists = {"evidence_index": d["evidence_index"], "related_claims": d["related_claims"],
                 "restatement_log": d["restatement_log"], **d["gap_report"]}
        for t, rows in lists.items():
            for r in rows:
                seen.setdefault(t, set()).update(k for k, v in r.items() if isinstance(v, dict) and "rows" in v)
        assert list(d["gap_report"]) == list(SECTIONS)
    assert seen["related_claims"] == seen["evidence_index"]
    assert {t: tuple(sorted(v)) for t, v in seen.items() if v and t != "related_claims"} == {
        t: tuple(sorted(v)) for t, v in rs.SUBS.items()}
    covered = {t for ts in SHEETS.values() for t in ts}
    assert set(dataset.COLUMNS) - {"related_claims"} <= covered and "pair_sides" in TITLES
    d = outs["complete", "all"][0]
    for x in REPORTS:
        assert [t["id"] for t in rs.tables(d, x)] == SHEETS[x]


@pytest.mark.parametrize("fmt", ["xlsx", "csv"])
def test_each_output_holds_exactly_its_tables_with_their_rows(outs, fmt):
    for (mode, name), (d, folder, _) in outs.items():
        for x in REPORTS:
            path = folder / (f"{stem(x)}.xlsx" if fmt == "xlsx" else stem(x))
            if fmt == "xlsx":
                titles = load_workbook(path, read_only=True).sheetnames
                assert titles == [TITLES["about"], *(TITLES[t] for t in SHEETS[x])]
                assert len({t.casefold() for t in titles}) == len(titles) and "history" not in {t.casefold() for t in titles}
            else:
                assert sorted(p.name for p in path.iterdir()) == sorted(f"{t}.csv" for t in ["about", *SHEETS[x]])
            got, _ = _tables_of(path)
            for t, cols, rows in expected(d, x):
                headings, data = got[t]
                assert headings == [c["heading"] for c in cols], (mode, name, x, t)
                assert data == [row_text(r, cols, fmt) for r in rows], (mode, name, x, t)
    # the sides: every claim_a and claim_b once, with its own row's values; related claims among them
    d = outs["complete", "docs"][0]
    sides = {t: rows for t, _, rows in expected(d, "restatement_log")}["pair_sides"]
    named = {c for p in d["restatement_log"] for c in (p["claim_a"], p["claim_b"])}
    assert sorted(r["id"] for r in sides) == sorted(named) and {r["id"] for r in d["related_claims"]} & named
    assert d["columns"]["evidence_index"] == d["columns"]["related_claims"]
    assert outs["sanitised", "docs"][0]["columns"]["evidence_index"] == outs["sanitised", "docs"][0]["columns"][
        "related_claims"]


# AC 3: values, through both readers

def test_every_xlsx_cell_through_openpyxl_and_numbers_are_numbers(outs):
    kinds = set()
    for (mode, name), (d, folder, _) in outs.items():
        for x in REPORTS:
            wb = load_workbook(folder / f"{stem(x)}.xlsx")
            for t, cols, rows in expected(d, x):
                ws = wb[TITLES[t]]
                values = list(ws.iter_rows(values_only=True))
                blank = next(i for i, r in enumerate(values) if all(v is None for v in r))
                assert ws.freeze_panes == f"A{blank + 3}"
                assert ws.auto_filter.ref.startswith(f"A{blank + 2}:") and ws.auto_filter.ref.endswith(
                    str(blank + 2 + len(rows)))
                for r, line in zip(rows, values[blank + 2:]):
                    for c, v in zip(cols, line):
                        want, raw = printed(rs.dotted(r, c["key"]), r, c["key"], "xlsx"), rs.dotted(r, c["key"])
                        if isinstance(raw, (int, float)) and not isinstance(raw, bool) and c["key"] not in r["withheld"]:
                            assert type(v) in (int, float) and "%.16g" % v == want, (t, c["key"], v)
                            kinds.add("number")
                        else:
                            assert ("" if v is None else v) == want, (t, c["key"], v, want)
                        kinds.add(type(raw).__name__)
    assert {"number", "bool", "list", "NoneType", "str"} <= kinds


def test_each_value_rule_is_exercised(outs):
    d = outs["sanitised", "docs"][0]
    plain = next(r for r in d["evidence_index"] if r["correction"] is None)
    assert "correction.rationale" in plain["withheld"]  # B5: withheld wins over the null object
    assert printed(None, plain, "correction.rationale", "xlsx") == WORDS["withheld"] == "withheld"
    complete = outs["complete", "docs"][0]
    pair = next(p for p in complete["restatement_log"] if any(p["outside_selection"]))
    a, b = (printed(rs.dotted(pair, f"outside_selection.{s}"), pair, f"outside_selection.{s}", "csv") for s in "ab")
    assert {a, b} == {WORDS["outside_selection"], ""}
    corrected = next(r for r in outs["complete", "all"][0]["evidence_index"] if r["correction"])
    assert printed(rs.dotted(corrected, "correction.created_at"), corrected, "correction.created_at", "csv") == \
        corrected["correction"]["created_at"]
    nothing = next(r for r in complete["evidence_index"] if r["correction"] is None)
    assert printed(rs.dotted(nothing, "correction.created_at"), nothing, "correction.created_at", "xlsx") == ""
    assert printed(["a", "b"], {}, "needs_renewed_review", "csv") == "a; b" and printed([3, 4], {}, "claims", "xlsx") == "3; 4"
    assert printed(True, {}, "binds", "xlsx") == "yes" and printed(False, {}, "binds", "csv") == "no"
    assert printed(1260.0, {}, "value", "xlsx") == "1260" and printed(1260.0, {}, "value", "csv") == "1260.0"
    assert printed("", {}, "note", "xlsx") == printed(None, {}, "note", "csv") == ""
    labelled = [c["key"] for c in complete["columns"]["evidence_index"]]
    assert labelled[-1] == "label" and "label" not in [c["key"] for c in d["columns"]["evidence_index"]]
    assert d["words"] == WORDS


def test_csv_numbers_are_bare_and_text_is_quoted(outs):
    """QUOTE_NONNUMERIC on read: an unquoted field becomes a float (and an unquoted text would fail), so the read
    tells quoted from bare exactly."""
    d, folder, _ = outs["complete", "all"]
    for x in REPORTS:
        for t, cols, rows in expected(d, x):
            with open(folder / stem(x) / f"{t}.csv", encoding="utf-8-sig", newline="") as fh:
                got = list(csv.reader(fh, csv.excel, quoting=csv.QUOTE_NONNUMERIC))
            assert all(isinstance(v, str) for v in got[0])
            for r, line in zip(rows, got[1:]):
                for c, v in zip(cols, line):
                    raw = rs.dotted(r, c["key"])
                    if isinstance(raw, (int, float)) and not isinstance(raw, bool):
                        assert isinstance(v, float) and v == raw, (t, c["key"])
                    else:
                        assert isinstance(v, str), (t, c["key"], v)


# AC 4: read_back

def test_read_back_has_the_shared_shape(outs):
    for (mode, name), (d, folder, _) in outs.items():
        for x in REPORTS:
            for path, fmt in ((folder / f"{stem(x)}.xlsx", "xlsx"), (folder / stem(x), "csv")):
                back = read_back(path)
                assert back["data_version"] == {k: printed(v, {}, k, fmt) for k, v in d["data_version"].items()}
                want = {"evidence_index": ["evidence_index"], "restatement_log": ["restatement_log"],
                        "gap_report": [f"gap_report.{s}" for s in SECTIONS]}[x]
                assert list(back["rows"]) == want
                for key, rows in back["rows"].items():
                    t = key.split(".")[-1]
                    src = d[t] if t in d else d["gap_report"][t]
                    assert rows == [{c["key"]: printed(rs.dotted(r, c["key"]), r, c["key"], fmt)
                                     for c in d["columns"][t]} for r in src], (mode, name, path, key)
    assert set(dataset.COLUMNS) >= {"evidence_index", "related_claims"} and "pair_sides" not in dataset.COLUMNS


# AC 5, 6, 7: version, notices, words

def test_every_sheet_carries_the_version_notices_and_the_self_declared_rule(outs):
    for (mode, name), (d, folder, _) in outs.items():
        dv, fixed = d["data_version"], WORDS["fixed"]
        for x in REPORTS:
            tables = expected(d, x)
            any_attributed = any(is_attributed(rows, cols) for _, cols, rows in tables)
            book = dict(sheet._xlsx(folder / f"{stem(x)}.xlsx"))
            for t, cols, rows in tables:
                band, _, _ = _split(book[TITLES[t]])
                assert band[0][:7] == [fixed["identity"]["version"], str(dv["version"]), fixed["identity"]["built_at"],
                                       dv["built_at"], fixed["identity"]["sha256"], dv["sha256"],
                                       fixed["mode"][d["mode"]]]
                lines = [r[0] for r in band[1:]]
                want = [*d["notices"], *([WORDS["self_declared"]] if is_attributed(rows, cols) else [])]
                assert lines == want, (mode, name, t)
                assert all(r[1:] == [""] * (len(r) - 1) for r in band[1:])
                csv_data = csv_rows(folder / stem(x) / f"{t}.csv")
                assert all(r[0] == str(dv["version"]) for r in csv_data[1:])  # the data_version column on every row
            for about in (book[TITLES["about"]], csv_rows(folder / stem(x) / "about.csv")):
                width = max(len(r) for r in about)
                assert about[0] == fixed["about_csv"] + [""] * (width - 2)
                ident = [r[:2] for r in about[1:1 + len(dv)]]
                assert ident == [[fixed["identity"][k], str(v)] for k, v in dv.items()]
                assert ("History mark" in {r[0] for r in about}) == (mode == "complete")
                rest = [r[:2] for r in about[1 + len(dv):]]
                assert rest[:4] == [[fixed["dataset"]["built_at"], d["built_at"]],
                                    [fixed["dataset"]["dataset_schema"], str(d["dataset_schema"])],
                                    [fixed["dataset"]["gct_version"], d["gct_version"]], ["", fixed["mode"][d["mode"]]]]
                assert rest[4] == [fixed["dataset"]["include"], "; ".join(d["include"])]
                assert rest[5][0] == fixed["dataset"]["filter"]
                assert rest[6:] == [["", n] for n in d["notices"]] + (
                    [["", WORDS["self_declared"]]] if any_attributed else [])
    assert outs["sanitised", "all"][0]["notices"] == [dataset.PRIVACY_NOTICE, dataset.ARCHIVE_NOTICE]


def _strip(d, *, corrections=True, records=True, pairs_restatements=True, adjudications=True):
    """A copy of a dataset with chosen attributed rows removed (the renderer reads the dict, so the cases are cut
    from a built dataset)."""
    d = copy.deepcopy(d)
    claims = d["evidence_index"] + d["related_claims"]
    for r in claims:
        if corrections:
            r["correction"] = None
        if records:
            r["records"] = {**r["records"], "rows": []}
    for p in d["restatement_log"] + d["gap_report"]["candidates"]:
        if pairs_restatements:
            p["restatements"] = {**p["restatements"], "rows": []}
        if adjudications:
            p["adjudications"] = {**p["adjudications"], "rows": []}
    if pairs_restatements:
        d["unpaired_restatements"] = {**d["unpaired_restatements"], "rows": []}
    return d


def test_the_self_declared_line_follows_the_attributed_rows(outs, tmp_path):
    """B9: complete mode, since a sanitised correction.contributor is withheld."""
    d = outs["complete", "all"][0]
    cases = {"correction only": (_strip(d, corrections=False), "evidence_index", True),
             "adjudications only": (_strip(d, adjudications=False), "restatement_log", True),
             "none": (_strip(d), "evidence_index", False), "none, log": (_strip(d), "restatement_log", False)}
    assert any(r["correction"] for r in cases["correction only"][0]["evidence_index"])
    for case, (cut, report, want) in cases.items():
        folder = tmp_path / case.replace(" ", "-").replace(",", "")
        render(cut, folder, [report])
        book = dict(sheet._xlsx(folder / f"{stem(report)}.xlsx"))
        about = [r[1] for r in book[TITLES["about"]]] + [r[1] for r in csv_rows(folder / stem(report) / "about.csv")]
        assert about.count(WORDS["self_declared"]) == 2 * want, case
        marked = {t for t, cols, rows in expected(cut, report) if is_attributed(rows, cols)}
        assert bool(marked) == want
        for t in SHEETS[report]:
            assert (WORDS["self_declared"] in {r[0] for r in _split(book[TITLES[t]])[0]}) == (t in marked), (case, t)
        assert all(WORDS["self_declared"] not in json.dumps(csv_rows(p)) for p in (folder / stem(report)).iterdir()
                   if p.name != "about.csv")


def test_printed_words_never_codes_and_the_guard_passes(outs):
    codes = set(adjudicate.STATES) | {"reviewer_confirmed_issue", "unresolved_conflict", "explained_difference"}
    authored = labels.CONTRIBUTOR_TEXT | labels.IDENTITY | labels.DOCUMENT_TEXT | {"contributor", "company", "text"}
    states = re.compile(r"^[a-z]+(_[a-z]+)+$")
    for (mode, name), (d, folder, _) in outs.items():
        for x in REPORTS:
            book = dict(sheet._xlsx(folder / f"{stem(x)}.xlsx"))
            for title in book:
                assert not GUARD.search(title)
            for t, cols, rows in expected(d, x):
                band, headings, data = _split(book[TITLES[t]])
                assert not GUARD.search(json.dumps([band, headings]))
                for line in data:
                    for c, v in zip(cols, line):
                        assert v not in codes, (t, c["key"], v)
                        if c["key"].endswith(("state", "state_text", "kind", "outcome", "decision")):
                            assert not states.match(v), (t, c["key"], v)
                        if c["key"] not in authored and not c["key"].endswith(tuple("." + k for k in authored)):
                            assert not GUARD.search(v), (t, c["key"], v)
            assert not GUARD.search(json.dumps([r[0] for r in book[TITLES["about"]]]))


# AC 8, 9: cell safety

SEEDS = ['=HYPERLINK("http://example.invalid","GCT-FORMULA-CANARY")', "+1+1", "-2+3", "@SUM(A1)", "a;=1+1", "#N/A",
         "", "\tleading tab", "\rleading return", "x\r\ny", "nu\x00l", "be\x07ll", "non\U0000FFFEchar", "non\U0000FFFFchar",
         "-12.5", "-1,234", "\t-1"]
NAMES = ("=x", "-x")  # two shareable documents' names


def step1(s):
    return re.sub("[\x00-\x08\x0b\x0c\x0e-\x1f\U0000FFFE\U0000FFFF]", "\N{REPLACEMENT CHARACTER}", s)


def as_xlsx(s):
    return step1(s).replace("\r\n", "\n").replace("\r", "\n")


def as_csv(s):
    s = step1(s)
    if s[:1] in ("\t", "\r") or (s[:1] in ("=", "+", "-", "@") and not re.fullmatch(r"[+-]?\d[\d,]*(\.\d+)?", s)):
        return "'" + s
    return s


def seeded(d, field):
    """A copy of a built dataset with one new evidence index row per seed: the seed in field (a document text, so it
    ships in sanitised mode) or, for "conclusion", one new review record per seed under the first claim; and two
    documents renamed =x and -x."""
    d = copy.deepcopy(d)
    first = d["evidence_index"][0]
    if field == "conclusion":
        base = copy.deepcopy(next(x for r in d["evidence_index"] for x in r["records"]["rows"]))
        first["records"]["rows"] += [{**base, "id": 900000 + i, "claim_id": first["id"], "conclusion": s}
                                     for i, s in enumerate(SEEDS)]
    else:
        d["evidence_index"] += [{**first, "id": 900000 + i, "claim_id": 900000 + i, field: s,
                                 "records": {**first["records"], "rows": []},
                                 "disagreements": {**first["disagreements"], "rows": []}} for i, s in enumerate(SEEDS)]
    for doc, name in zip(d["documents"], NAMES):
        doc["name"] = name
    return d


@pytest.fixture(scope="module")
def safety(outs, tmp_path_factory):
    """{mode: (dataset, folder, writes)}: complete mode with the seeds as contributor text, sanitised as document
    text."""
    got = {}
    for mode, field in (("complete", "conclusion"), ("sanitised", "metric_text")):
        d = seeded(outs[mode, "all"][0], field)
        folder = tmp_path_factory.mktemp(f"safety-{mode}")
        got[mode] = d, folder, render(d, folder, ["evidence_index"])
    return got


SEEDED = {"complete": ("records", "Conclusion", "ID"), "sanitised": ("evidence_index", "Metric as printed", "Claim")}


def test_xlsx_cells_are_inline_strings_never_formulas_or_errors(safety):
    for mode, (d, folder, _) in safety.items():
        path = folder / "evidence-index.xlsx"
        for name, xml in worksheets(path).items():
            assert b"<f>" not in xml and b"<f " not in xml, name
            for c in ElementTree.fromstring(xml).iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}c"):
                kids = [k.tag.split("}")[1] for k in c]
                assert (c.get("t"), kids) in (("inlineStr", ["is"]), ("n", ["v"])), (name, c.attrib, kids)
        with zipfile.ZipFile(path) as z:
            assert not [n for n in z.namelist() if n.startswith("xl/media/")]
        table, heading, ids = SEEDED[mode]
        _, headings, data = _split(dict(sheet._xlsx(path))[TITLES[table]])
        got = [r[headings.index(heading)] for r in data if r[headings.index(ids)].startswith("9000")]
        assert got == [as_xlsx(s) for s in SEEDS], mode
        ws = load_workbook(path)[TITLES[table]]
        cells = [r[headings.index(heading)] for r in ws.iter_rows(values_only=True)
                 if isinstance(r[headings.index(ids)], int) and r[headings.index(ids)] >= 900000]
        assert [("" if v is None else v) for v in cells] == [as_xlsx(s) for s in SEEDS]
        docs = _split(dict(sheet._xlsx(path))[TITLES["documents"]])
        assert [r[docs[1].index("Name")] for r in docs[2]][:2] == list(NAMES)


def test_csv_fields_are_prefixed_quoted_and_replaced(safety, outs):
    for mode, (d, folder, out) in safety.items():
        table, heading, ids = SEEDED[mode]
        rows = csv_rows(folder / "evidence-index" / f"{table}.csv")
        got = [r[rows[0].index(heading)] for r in rows[1:] if r[rows[0].index(ids)].startswith("9000")]
        assert got == [as_csv(s) for s in SEEDS], mode
        assert got[SEEDS.index("-12.5")] == "-12.5" and got[SEEDS.index("-1,234")] == "-1,234"
        assert got[SEEDS.index("\t-1")] == "'\t-1" and got[SEEDS.index("x\r\ny")] == "x\r\ny"
        assert sum(g.startswith("'") for g in got) == 7
        docs = csv_rows(folder / "evidence-index" / "documents.csv")
        assert [r[docs[0].index("Name")] for r in docs[1:]][:2] == ["'=x", "'-x"]
        # the notes count by hand: four replaced characters; seven seeds and two names prefixed; none unseeded
        assert notes_counts(outs[mode, "all"][2]) == {"xlsx": [0], "csv": [0, 0]}
        assert notes_counts(out) == {"xlsx": [4], "csv": [4, 9]}
        assert out["xlsx"]["notes"] == ["xlsx: 4 characters replaced with U+FFFD (control characters, U+FFFE and"
                                        " U+FFFF)."]


# AC 10: refusals before any file

def test_a_cell_or_a_sheet_too_big_for_xlsx_is_refused_before_any_file(outs, tmp_path, monkeypatch):
    d = outs["complete", "all"][0]
    record = next(x for r in d["evidence_index"] for x in r["records"]["rows"])
    for i, text in enumerate(("a" * 32_768, "\U0001F600" * 16_384)):
        cut = copy.deepcopy(d)
        rec = next(x for r in cut["evidence_index"] for x in r["records"]["rows"] if x["id"] == record["id"])
        rec["conclusion"] = text
        folder = tmp_path / f"long-{i}"
        folder.mkdir()
        with pytest.raises(rs.Refusal) as e:
            rs.write_xlsx(cut, list(REPORTS), folder, {})
        assert e.value.cause == (f"The Review records table of the evidence-index report has a value longer than an"
                                 f" XLSX cell can hold (row {record['id']}, column Conclusion).")
        assert e.value.next_step == "Export CSV or JSON for this selection, or narrow the filter."
        assert list(folder.iterdir()) == []
        rs.write_csv(cut, list(REPORTS), folder, {})  # CSV has no such limit
        assert text in [x for r in csv_rows(folder / "evidence-index" / "records.csv") for x in r]
    ok = copy.deepcopy(d)  # 32,767 units is the limit itself, not over it
    next(x for r in ok["evidence_index"] for x in r["records"]["rows"])["conclusion"] = "a" * 32_767
    (tmp_path / "ok").mkdir()
    rs.write_xlsx(ok, ["evidence_index"], tmp_path / "ok", {})
    # a cell only the third report (the gap report) prints: the first two reports' files are not written either (R3.1)
    third = copy.deepcopy(d)
    third["gap_report"]["extraction_limits"][0]["reason"] = "a" * 32_768
    (tmp_path / "third").mkdir()
    with pytest.raises(rs.Refusal) as e:
        rs.write_xlsx(third, list(REPORTS), tmp_path / "third", {})
    assert e.value.cause.startswith("The Extraction limits table of the gap-report report has a value longer")
    assert list((tmp_path / "third").iterdir()) == []
    # a sheet over the row limit
    rows = 1 + len(d["notices"]) + 1 + 1 + len(d["documents"])
    monkeypatch.setattr(rs, "MAX_ROWS", rows - 1)
    (tmp_path / "rows").mkdir()
    with pytest.raises(rs.Refusal) as e:
        rs.write_xlsx(d, list(REPORTS), tmp_path / "rows", {})
    assert "more rows than an XLSX sheet can hold" in e.value.cause and list((tmp_path / "rows").iterdir()) == []


def test_the_refusal_reaches_the_user_as_the_commands_refusal(ws, gct, monkeypatch):
    names = internal_names(ws)
    monkeypatch.setattr(rs, "MAX_ROWS", 25)  # above the About sheet's rows, below the evidence index sheet's
    before = sorted((ws / "exports").rglob("*")) if (ws / "exports").exists() else []
    human = gct("report", "all", "--format", "csv,xlsx", "--workspace", ws)
    machine = gct("report", "all", "--format", "xlsx", "--workspace", ws, "--json")
    assert human.code == machine.code == layout.REFUSED and human.out == ""
    cause = machine.data["cause"]
    assert cause == "The Evidence index table of the evidence-index report has more rows than an XLSX sheet can hold (25)."
    assert human.err == f"error: {cause}\nnext: {rs.NEXT_STEP}\n" and machine.data["next_action"] == rs.NEXT_STEP
    for text in (human.err, machine.out):
        for word in [*INTERNAL_WORDS, *names]:
            assert word not in text, word
    assert sorted((ws / "exports").rglob("*")) == before
    assert gct("report", "all", "--format", "csv", "--workspace", ws).code == 0


# AC 11: encoding and empty tables

def test_csv_encoding_and_empty_tables_keep_their_headings(w, env, tmp_path):
    d, _ = build(w, mode="complete", filter={"document": [5]})  # the deleted document: no claim
    assert d["evidence_index"] == []
    render(d, tmp_path / "out")
    for x in REPORTS:
        names = sorted(p.name for p in (tmp_path / "out" / stem(x)).iterdir())
        assert names == sorted(f"{t}.csv" for t in ["about", *SHEETS[x]])
        for p in (tmp_path / "out" / stem(x)).iterdir():
            raw = p.read_bytes()
            assert raw.startswith(b"\xef\xbb\xbf\"") and raw.endswith(b"\r\n") and b"\n" not in raw.replace(b"\r\n", b"")
    assert csv_rows(tmp_path / "out" / "evidence-index" / "records.csv") == [
        [c["heading"] for c in d["columns"]["records"]]]
    book = dict(sheet._xlsx(tmp_path / "out" / "evidence-index.xlsx"))
    assert _split(book[TITLES["records"]])[1:] == ([c["heading"] for c in d["columns"]["records"]], [])
    assert _split(book[TITLES["evidence_index"]])[0][1][0] == d["notices"][0]


# AC 12: privacy

def printed_strings(d) -> str:
    """Every string a report output prints: table cells, notices and About values."""
    out = [*d["notices"], *map(str, d["data_version"].values()), json.dumps(d["filter"], ensure_ascii=False),
           *d["include"], d["built_at"], d["gct_version"]]
    for x in REPORTS:
        for _, cols, rows in expected(d, x):
            out += [v for r in rows for v in row_text(r, cols, "csv")]
    return "\n".join(out)


def test_the_complete_output_holds_its_printed_seeds_and_the_sanitised_none(w, env, tmp_path):
    complete_only = [s for s in f.SEEDS if s["expect"] == "complete_only"]
    never = [s for s in f.SEEDS if s["expect"] == "never"]
    flt = {"document": [8, 1, 19], "company": [1, 2, 3], "text": f.METRIC}
    for case, kw in (("whole", {}), ("filter", {"filter": flt})):
        d, _ = build(w, mode="complete", **kw)
        render(d, tmp_path / case / "complete")
        text = printed_strings(d)
        want = {s["name"] for s in complete_only if s["kind"] in ("document_text", "contributor") and s["value"] in text}
        assert set(found(tmp_path / case / "complete", complete_only)) == want, case
        if case == "whole":  # the expected set holds both kinds, so the check can fail
            assert {"document_text", "contributor"} <= {s["kind"] for s in complete_only if s["name"] in want}
        assert found(tmp_path / case / "complete", never) == {}
        s, _ = build(w, **kw)
        render(s, tmp_path / case / "sanitised")
        assert found(tmp_path / case / "sanitised", f.SEEDS) == {}
        for mode in ("complete", "sanitised"):
            for x in REPORTS:
                with zipfile.ZipFile(tmp_path / case / mode / f"{stem(x)}.xlsx") as z:
                    core = z.read("docProps/core.xml")
                    assert b"<dc:creator>openpyxl</dc:creator>" in core
                    assert not re.search(rb"/(Users|home)/", b"".join(z.read(n) for n in z.namelist()))
                    assert not [n for n in z.namelist() if n.startswith("xl/media/")]
        if case == "filter":  # K11: the canary text in the complete echo only; the sanitised echo withholds it
            private = max(flt["document"])  # the fixture's private PDF has the highest id
            assert private not in {x["id"] for x in s["documents"]}
            for mode, dd in (("complete", d), ("sanitised", s)):
                for about in (dict(sheet._xlsx(tmp_path / case / mode / "evidence-index.xlsx"))[TITLES["about"]],
                              csv_rows(tmp_path / case / mode / "evidence-index" / "about.csv")):
                    echo = json.loads(next(r[1] for r in about if r[0] == WORDS["fixed"]["dataset"]["filter"]))
                    if mode == "complete":
                        assert echo == dd["filter"] and echo["text"] == f.METRIC and private in echo["document"]
                    else:
                        assert echo == {**{k: v for k, v in dd["filter"].items() if k != "withheld"},
                                        "text": WORDS["withheld"]}
                        assert set(echo["document"]) <= {x["id"] for x in s["documents"]}
            assert "DOC-METRIC-01" in want


def test_sanitised_contributors_are_pseudonyms(outs):
    """C7, with no inclusion class (B9: with --include identity the stored names would ship by opt-in)."""
    for name in ("all", "docs"):
        d, folder, _ = outs["sanitised", name]
        seen = []
        for x in REPORTS:
            for t, cols, rows in expected(d, x):
                if "contributor" in [c["key"] for c in cols]:
                    data = csv_rows(folder / stem(x) / f"{t}.csv")
                    seen += [r[data[0].index("Contributor")] for r in data[1:]]
        assert seen and all(re.fullmatch(r"Contributor \d+", v) for v in seen)


# AC 13, 14: behaviour

def test_archive_missing_states_and_non_binding_records(outs):
    d, folder, _ = outs["complete", "all"]
    book = dict(sheet._xlsx(folder / "gap-report.xlsx"))
    archived = _split(book[TITLES["archived"]])[2]
    assert len(archived) == len(d["gap_report"]["archived"]) == 2
    assert any("archived documents" in n for n in d["notices"])  # the F7 line, on every sheet (notices test)
    expectations = _split(book[TITLES["expectations"]])
    states = {r[expectations[1].index("State")] for r in expectations[2]}
    assert states and states <= {missing.INCOMPLETE, missing.NOT_FOUND_IN, missing.CONFIRMED}
    assert _split(book[TITLES["reasons"]])[2] and _split(book[TITLES["declared"]])[2]
    limits = _split(book[TITLES["extraction_limits"]])
    assert missing.CHARTS in {r[limits[1].index("Reason")] for r in limits[2]}
    records = _split(dict(sheet._xlsx(folder / "evidence-index.xlsx"))[TITLES["records"]])
    earlier = [r for r in records[2] if r[records[1].index("Binds")] == "no"]
    assert earlier and {r[records[1].index("Version note")] for r in earlier} == {WORDS["earlier_version"]}
    assert all(r[records[1].index("On claim version")] for r in earlier)
    s, sfolder, _ = outs["sanitised", "all"]
    sbook = dict(sheet._xlsx(sfolder / "gap-report.xlsx"))
    assert _split(sbook[TITLES["archived"]])[2] == _split(sbook[TITLES["deleted"]])[2] == []
    for x in REPORTS:
        for title, rows in sheet._xlsx(sfolder / f"{stem(x)}.xlsx"):
            assert [dataset.ARCHIVE_NOTICE] in [r[:1] for r in rows] if title != TITLES["about"] else [
                "", dataset.ARCHIVE_NOTICE] in [r[:2] for r in rows], title
        assert ["", dataset.ARCHIVE_NOTICE] in [r[:2] for r in csv_rows(sfolder / stem(x) / "about.csv")]


def test_a_correction_after_a_review_shows_needs_renewed_review_joined(ws, gct):
    """The fixture holds no renewal entry, so a correction on a claim with a binding review makes one (E6-D8)."""
    d, _ = build(ws, mode="complete")
    claim = next(r for r in d["evidence_index"] if any(x["binds"] and x["kind"] == "auditor review"
                                                         for x in r["records"]["rows"]))
    assert gct("correct", claim["claim_id"], "--value", "1,111", *WHO, "--workspace", ws).code == 0
    assert gct("publish", "build", "--workspace", ws).code == 0
    r = gct("report", "evidence-index", "--format", "xlsx,csv", "--complete", "--workspace", ws, "--json")
    folder = Path(r.data["folder"])
    d, _ = build(ws, mode="complete")
    joined = [WORDS["fixed"]["separator"].join(r["needs_renewed_review"]) for r in d["evidence_index"]]
    assert any(joined) and dataset.HEADINGS["needs_renewed_review"] == "Needs renewed review"
    for fmt, path in (("xlsx", folder / "evidence-index.xlsx"), ("csv", folder / "evidence-index")):
        headings, data = _tables_of(path)[0]["evidence_index"]
        assert [x[headings.index("Needs renewed review")] for x in data] == [printed(j, {}, "x", fmt) for j in joined]
        records = _tables_of(path)[0]["records"]
        earlier = [x for x in records[1] if x[records[0].index("Binds")] == "no"]
        assert earlier and {x[records[0].index("Version note")] for x in earlier} == {WORDS["earlier_version"]}


def test_read_back_reads_a_csv_field_above_the_csv_modules_default_limit(ws, gct):
    """CSV has no cell limit, so read_back reads a field longer than csv's default 131,072 characters, whole, and
    leaves the process limit as it was (Codex review of 5f7746b)."""
    d, _ = build(ws, mode="complete")
    claim = d["evidence_index"][0]["claim_id"]
    long = "r" * 131_072 + "!"
    assert gct("correct", claim, "--value", "1,111", "--as", "ana", "--capacity", "analyst", "--rationale", long,
               "--workspace", ws).code == 0
    assert gct("publish", "build", "--workspace", ws).code == 0
    r = gct("report", "evidence-index", "--format", "json,csv", "--complete", "--workspace", ws, "--json")
    assert r.code == 0, r.out
    folder = Path(r.data["folder"])
    exported = json.loads((folder / "dataset.json").read_text(encoding="utf-8"))
    want = [x["correction"] and x["correction"]["rationale"] for x in exported["evidence_index"]]
    assert long in want
    limit = csv.field_size_limit()
    rows = read_back(folder / "evidence-index")["rows"]["evidence_index"]
    assert [x["correction.rationale"] for x in rows] == [w or "" for w in want]
    assert csv.field_size_limit() == limit


def test_a_deletion_shows_the_new_sha256_under_the_same_number(ws, gct):
    def export_(*argv):
        r = gct("report", "all", "--format", "xlsx,csv", "--complete", *argv, "--workspace", ws, "--json")
        assert r.code == 0, r.out
        return Path(r.data["folder"])

    first = export_()
    kept = {p: p.read_bytes() for p in first.rglob("*") if p.is_file()}
    token = sqlite3.connect(ws / layout.DB_NAME).execute("SELECT sha256 FROM document WHERE id = 1").fetchone()[0][:12]
    assert gct("doc", "delete", 1, "--confirm", token, *WHO, "--workspace", ws).code == 0
    after = export_()
    assert {p: p.read_bytes() for p in first.rglob("*") if p.is_file()} == kept  # earlier files stay
    for path in ("evidence-index.xlsx", "evidence-index"):
        a, b = read_back(first / path)["data_version"], read_back(after / path)["data_version"]
        assert a["version"] == b["version"] and a["sha256"] != b["sha256"]
    rows = read_back(after / "evidence-index")["rows"]["evidence_index"]
    assert rows and all(r["document_id"] != "1" for r in rows)


def test_every_document_private_writes_headings_only(ws, gct):
    for x in gct("doc", "list", "--workspace", ws, "--json").data["documents"]:
        if x["state"] != "deleted" and x["privacy"] == "shareable":
            assert gct("doc", "label", x["id"], "--privacy", "private", *WHO, "--workspace", ws).code == 0
    assert gct("publish", "build", "--workspace", ws).code == 0
    out = {}
    for argv in ((), ("--complete",)):
        r = gct("report", "all", "--format", "xlsx,csv", *argv, "--workspace", ws, "--json")
        assert r.code == 0, r.out
        out[argv] = Path(r.data["folder"])
    folder, v = out[()], read_back(out[()] / "evidence-index")["data_version"]["version"]
    notices = [dataset.sanitised_empty(int(v)), dataset.PRIVACY_NOTICE, dataset.ARCHIVE_NOTICE]
    for x in REPORTS:
        for t, (headings, data) in _tables_of(folder / f"{stem(x)}.xlsx")[0].items():
            assert data == [], t
        for title, rows in sheet._xlsx(folder / f"{stem(x)}.xlsx"):
            if title != TITLES["about"]:
                assert [r[0] for r in _split(rows)[0][1:]] == notices
        assert [r[1] for r in csv_rows(folder / stem(x) / "about.csv") if r[0] == ""][1:] == notices
        for t, (headings, data) in _tables_of(folder / stem(x))[0].items():
            assert data == [], t
    assert found(folder, f.SEEDS) == {}
    docs = gct("doc", "list", "--workspace", ws, "--json").data["documents"]
    names = [x["logical_name"] for x in docs if x["state"] != "deleted"]
    for p in folder.rglob("*"):
        for n in names:
            assert n not in p.name
    blob = "".join(json.dumps(sheet._xlsx(p)) for p in folder.glob("*.xlsx"))
    blob += "".join(p.read_text(encoding="utf-8-sig") for p in folder.rglob("*.csv"))
    assert not [n for n in names if n in blob]
    complete = "".join(json.dumps(sheet._xlsx(p)) for p in out[("--complete",)].glob("*.xlsx"))
    assert all(n in complete for n in names)  # the complete output names each document, so the check can fail


# AC 15: scale

def test_twenty_thousand_claims_render_in_both_formats(outs, tmp_path):
    d = copy.deepcopy(outs["complete", "all"][0])
    first = d["evidence_index"][0]
    d["evidence_index"] = [{**first, "id": i, "claim_id": i, "records": {**first["records"], "rows": []},
                            "disagreements": {**first["disagreements"], "rows": []}} for i in range(1, 20_001)]
    start = time.monotonic()
    render(d, tmp_path, ["evidence_index"])
    took = time.monotonic() - start
    for path in (tmp_path / "evidence-index.xlsx", tmp_path / "evidence-index"):
        rows = read_back(path)["rows"]["evidence_index"]
        assert [r["claim_id"] for r in rows] == [str(i) for i in range(1, 20_001)]
    assert took < 120, took


# R3 preservation

def test_a_render_changes_no_workspace_file_and_csv_is_byte_stable(w, env, tmp_path):
    root = Path(shutil.copytree(w, tmp_path / "ws"))
    conn = store.connect(root / layout.DB_NAME)
    before = dump(root)
    published = {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in (root / "published").iterdir()}
    d, _ = dataset.build(conn, root, mode="complete")
    one, two = render(d, root / "exports" / "one"), render(d, root / "exports" / "two")
    assert dump(root) == before
    conn.close()
    assert {p.name: hashlib.sha256(p.read_bytes()).hexdigest() for p in (root / "published").iterdir()} == published
    assert one == two
    for name in one["csv"]["files"]:
        assert (root / "exports" / "one" / name).read_bytes() == (root / "exports" / "two" / name).read_bytes()
    for name in one["xlsx"]["files"]:
        assert sheet._xlsx(root / "exports" / "one" / name) == sheet._xlsx(root / "exports" / "two" / name)


def test_docs_describe_the_formats():
    text = (ROOT / "docs" / "reports.md").read_text(encoding="utf-8")
    section = " ".join(text.split("\n## XLSX and CSV\n", 1)[1].split("\n## ", 1)[0].split())
    for phrase in ("pair sides", "U+FFFD", "byte order mark", "32,767", rs.NEXT_STEP, "JSON export"):
        assert phrase in section, phrase
